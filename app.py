import os
import secrets
import hashlib
import calendar
import socket
import csv
import io
import json
import re
from datetime import datetime, timedelta, timezone
from functools import wraps

from flask import (
    Flask, render_template, redirect, url_for,
    request, flash, send_from_directory, abort,
    send_file, session, g
)
import openpyxl
import xlrd
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename

from config import Config
from db import init_db, get_conn, is_integrity_error, utcnow_iso


PAGE_VISIBILITY_ENDPOINTS = {
    "admin_panel": "Admin",
    "admin_permissions_page": "Admin Permissions",
    "admin_courses": "Courses",
    "admin_edit_course": "Course Editor",
    "admin_delete_course": "Courses",
    "admin_subjects": "Subjects",
    "admin_edit_subject": "Subject Editor",
    "admin_remove_subject_course": "Subjects",
    "admin_delete_subject": "Subjects",
    "admin_assign_subject": "Subjects",
    "admin_devices": "Device Management",
    "admin_device_settings": "Device Management",
    "admin_toggle_device": "Device Management",
    "admin_update_device_timer": "Device Management",
    "admin_delete_device": "Device Management",
    "admin_release_device": "Device Management",
    "admin_generate_lab_key": "Admin",
    "admin_deactivate_lab": "Admin",
    "admin_login_activity": "Login Activity",
    "admin_force_logout_session": "Login Activity",
    "admin_students": "Registration",
    "admin_logout_student_device": "Registration",
    "admin_import_students": "Registration",
    "admin_export_students": "Registration",
    "admin_edit_student": "Student Editor",
    "admin_delete_student": "Registration",
    "admin_deactivate_student": "Registration",
    "admin_activate_student": "Registration",
    "admin_documents": "Uploaded Documents",
    "admin_upload_document": "Uploaded Documents",
    "admin_show_document": "Document Viewer",
    "admin_update_document": "Document Editor",
    "admin_delete_document": "Uploaded Documents",
    "admin_attendance": "View Attendance",
    "admin_reports": "Reports",
    "admin_password_page": "Password",
    "admin_forgot_user_password": "Forgot User Password",
    "admin_forgot_password": "Admin Forgot Password",
    "dashboard": "Student Dashboard",
    "attendance": "Student Attendance",
    "documents": "Student Documents",
    "download_document": "Student Documents",
    "lab_entry": "Lab Entry",
    "register": "Student Registration",
    "forgot_password": "Forgot Password",
    "admin_view_attendance": "Session Attendance",
}

ADMIN_PAGE_PERMISSION_DEFINITIONS = (
    ("dashboard", "Dashboard", "/admin", "dashboard", "view"),
    ("courses", "Courses", "/admin/courses", "course_management", "view"),
    ("subjects", "Subjects", "/admin/subjects", "subject_management", "view"),
    ("devices", "Device Management", "/admin/devices", "lab_management", "view"),
    ("login_activity", "Login Activity", "/admin/login-activity", "user_management", "view"),
    ("registration", "Registration", "/admin/students", "student_management", "view"),
    ("documents", "Uploaded Documents", "/admin/documents", "documents", "view"),
    ("attendance", "View Attendance", "/admin/attendance", "attendance", "view"),
    ("reports", "Reports", "/admin/reports", "reports", "view"),
    ("password", "Password", "/admin/password", "settings", "edit"),
)

ADMIN_PAGE_PERMISSION_ENDPOINTS = {
    "admin_panel": "dashboard",
    "admin_generate_lab_key": "dashboard",
    "admin_deactivate_lab": "dashboard",
    "admin_courses": "courses",
    "admin_edit_course": "courses",
    "admin_delete_course": "courses",
    "admin_subjects": "subjects",
    "admin_edit_subject": "subjects",
    "admin_remove_subject_course": "subjects",
    "admin_delete_subject": "subjects",
    "admin_assign_subject": "subjects",
    "admin_devices": "devices",
    "admin_device_settings": "devices",
    "admin_toggle_device": "devices",
    "admin_update_device_timer": "devices",
    "admin_delete_device": "devices",
    "admin_release_device": "devices",
    "admin_login_activity": "login_activity",
    "admin_force_logout_session": "login_activity",
    "admin_students": "registration",
    "admin_logout_student_device": "registration",
    "admin_import_students": "registration",
    "admin_export_students": "registration",
    "admin_edit_student": "registration",
    "admin_delete_student": "registration",
    "admin_deactivate_student": "registration",
    "admin_activate_student": "registration",
    "admin_documents": "documents",
    "admin_upload_document": "documents",
    "admin_show_document": "documents",
    "admin_update_document": "documents",
    "admin_delete_document": "documents",
    "admin_attendance": "attendance",
    "admin_view_attendance": "attendance",
    "admin_reports": "reports",
    "admin_password_page": "password",
    "admin_forgot_user_password": "password",
    "admin_forgot_password": "password",
}


def get_page_visibility(db, page_name):
    row = db.execute(
        "SELECT is_visible FROM page_visibility WHERE page_name = ?", (page_name,)
    ).fetchone()
    return bool(row["is_visible"]) if row else True


def page_name_for_admin_endpoint(endpoint, args=None):
    if endpoint == "admin_export_attendance":
        return "Reports" if args and args.get("view") == "report" else "View Attendance"
    return PAGE_VISIBILITY_ENDPOINTS.get(endpoint)


def admin_page_key_for_endpoint(endpoint, args=None):
    if endpoint == "admin_export_attendance":
        return "reports" if args and args.get("view") == "report" else "attendance"
    return ADMIN_PAGE_PERMISSION_ENDPOINTS.get(endpoint)


def get_admin_page_permissions(db, admin_id):
    module_permissions = get_admin_permissions(db, admin_id)
    saved_rows = db.execute(
        "SELECT page_key, is_enabled FROM admin_page_permission WHERE admin_id = ?",
        (admin_id,),
    ).fetchall()
    saved_permissions = {row["page_key"]: bool(row["is_enabled"]) for row in saved_rows}
    return {
        page_key: saved_permissions.get(
            page_key,
            bool(module_permissions.get(module_name, {}).get(action_name, False)),
        )
        for page_key, _page_name, _page_url, module_name, action_name
        in ADMIN_PAGE_PERMISSION_DEFINITIONS
    }


class AnonymousUser:
    is_authenticated = False
    is_admin = False
    name = None


class CurrentUser:
    """Thin wrapper so templates can do current_user.name / .is_admin etc."""
    is_authenticated = True

    def __init__(self, row):
        self._row = row

    def __getattr__(self, item):
        try:
            return self._row[item]
        except (IndexError, KeyError):
            raise AttributeError(item)

    @property
    def is_admin(self):
        return self._row["role"] == "admin" or self.is_super_admin

    @property
    def is_super_admin(self):
        return bool(self._row["super_admin"])


def format_ist(value):
    if not value:
        return "-"
    try:
        raw_value = str(value).replace("T", " ")[:19]
        utc_value = datetime.strptime(raw_value, "%Y-%m-%d %H:%M:%S").replace(tzinfo=timezone.utc)
        ist_value = utc_value.astimezone(timezone(timedelta(hours=5, minutes=30)))
        return ist_value.strftime("%d-%m-%Y %I:%M:%S %p IST")
    except (TypeError, ValueError):
        return str(value)


def format_ist_date(value):
    formatted_value = format_ist(value)
    return formatted_value.split(" ", 1)[0] if formatted_value != "-" else formatted_value


def format_course_label(name, year):
    name = str(name or "").strip()
    year = str(year or "").strip()
    year_match = re.fullmatch(r"(\d+)\s*(?:YEAR)?", year, re.IGNORECASE)
    if year_match:
        year = f"{year_match.group(1)} YEAR"
    elif year:
        year = year.upper()
    return " ".join(value for value in (name, year) if value)


def normalize_course_reference(value):
    return " ".join(str(value or "").strip().lower().replace("_", " ").split())


def build_course_lookup(courses):
    lookup = {}
    for course in courses:
        name = str(course["name"] or "").strip()
        year = str(course["year"] or "").strip()
        aliases = set()
        if name:
            aliases.add(normalize_course_reference(name))
        if year:
            aliases.add(normalize_course_reference(year))
            year_digits = re.sub(r"\D+", "", year)
            if year_digits:
                aliases.update({
                    normalize_course_reference(year_digits),
                    normalize_course_reference(f"{year_digits} year"),
                    normalize_course_reference(f"{year_digits}YEAR"),
                    normalize_course_reference(f"{name} {year}"),
                    normalize_course_reference(f"{name} {year_digits}"),
                    normalize_course_reference(f"{name} {year_digits} year"),
                })
        label = format_course_label(name, year)
        normalized_label = normalize_course_reference(label)
        if normalized_label:
            aliases.add(normalized_label)
        if name:
            aliases.add(normalize_course_reference(name))
        for alias in aliases:
            if alias:
                lookup.setdefault(alias, []).append(course["id"])
    return lookup


def resolve_course_id(course_lookup, raw_course_value):
    normalized = normalize_course_reference(raw_course_value)
    if not normalized:
        return None
    if normalized in course_lookup:
        candidates = course_lookup[normalized]
        return candidates[0] if len(candidates) == 1 else candidates[0]
    if " " in normalized:
        base = normalized.rsplit(" ", 1)
        if len(base) == 2:
            yearless = base[0]
            if yearless in course_lookup:
                candidates = course_lookup[yearless]
                return candidates[0] if len(candidates) == 1 else candidates[0]
    return None


def calculate_validity_minutes(start_value, end_value):
    if not start_value or not end_value:
        return 0
    try:
        start_dt = datetime.strptime(str(start_value).replace("T", " ")[:19], "%Y-%m-%d %H:%M:%S")
        end_dt = datetime.strptime(str(end_value).replace("T", " ")[:19], "%Y-%m-%d %H:%M:%S")
        delta = end_dt - start_dt
        return max(0, int(delta.total_seconds() // 60))
    except (TypeError, ValueError):
        return 0


IMPORT_COLUMNS = ("Student ID", "Student Name", "Course", "Mobile", "Email")
MOBILE_PATTERN = re.compile(r"[0-9]{10}")
EMAIL_PATTERN = re.compile(
    r"[A-Za-z0-9!#$%&'*+/=?^_`{|}~-]+(?:\.[A-Za-z0-9!#$%&'*+/=?^_`{|}~-]+)*"
    r"@(?:[A-Za-z0-9](?:[A-Za-z0-9-]{0,61}[A-Za-z0-9])?\.)+"
    r"[A-Za-z]{2,63}"
)


def is_valid_email(email):
    if not isinstance(email, str):
        return False
    email = email.strip()
    if len(email) > 255 or not EMAIL_PATTERN.fullmatch(email):
        return False
    local_part, domain = email.rsplit("@", 1)
    return len(local_part) <= 64 and all(len(label) <= 63 for label in domain.split("."))


def is_valid_mobile_number(mobile_no):
    return MOBILE_PATTERN.fullmatch(mobile_no) is not None


def normalize_import_header(value):
    return " ".join(str(value or "").strip().lower().replace("_", " ").split())


def read_student_import(file_storage):
    filename = (file_storage.filename or "").lower()
    raw_data = file_storage.read()
    if filename.endswith(".csv"):
        rows = csv.reader(io.StringIO(raw_data.decode("utf-8-sig")))
        return [row for row in rows if any(value is not None and str(value).strip() for value in row)]
    if filename.endswith(".xlsx"):
        workbook = openpyxl.load_workbook(io.BytesIO(raw_data), read_only=True, data_only=True)
        worksheet = workbook.active
        values = [list(row) for row in worksheet.iter_rows(values_only=True)]
        workbook.close()
        return [row for row in values if any(value is not None and str(value).strip() for value in row)]
    if filename.endswith(".xls"):
        workbook = xlrd.open_workbook(file_contents=raw_data)
        worksheet = workbook.sheet_by_index(0)
        rows = [worksheet.row_values(row_index) for row_index in range(worksheet.nrows)]
        return [row for row in rows if any(value is not None and str(value).strip() for value in row)]
    raise ValueError("Only .xlsx, .xls, and .csv files are supported.")


def create_app(config_class=Config, db_path=None):
    app = Flask(__name__)
    app.config.from_object(config_class)
    if db_path:
        app.config["DB_PATH"] = db_path
    os.makedirs(app.config["UPLOAD_FOLDER"], exist_ok=True)
    app.jinja_env.filters["ist"] = format_ist
    app.jinja_env.filters["ist_date"] = format_ist_date
    app.jinja_env.filters["course_label"] = format_course_label

    @app.before_request
    def protect_state_changing_requests():
        if not app.config.get("CSRF_ENABLED", True) or request.method not in {"POST", "PUT", "PATCH", "DELETE"}:
            return
        expected_token = session.get("_csrf_token", "")
        supplied_token = request.form.get("csrf_token", "") or request.headers.get("X-CSRF-Token", "")
        if not expected_token or not secrets.compare_digest(expected_token, supplied_token):
            abort(400, description="Invalid or missing CSRF token.")

    @app.context_processor
    def inject_csrf_token():
        if "_csrf_token" not in session:
            session["_csrf_token"] = secrets.token_urlsafe(32)

        def admin_page_access(endpoint):
            user = get_current_user()
            if not user.is_authenticated or not user.is_admin:
                return False
            page_name = page_name_for_admin_endpoint(endpoint, request.args)
            if page_name and not get_page_visibility(get_db(), page_name):
                return False
            page_key = admin_page_key_for_endpoint(endpoint, request.args)
            if page_key and not user.is_super_admin:
                if not get_admin_page_permissions(get_db(), user.id).get(page_key, False):
                    return False
            required = required_admin_permission(endpoint, "GET")
            if not required:
                return False
            return admin_can_access_module(user, required[0], required[1], get_db())

        def page_is_visible(endpoint):
            page_name = page_name_for_admin_endpoint(endpoint, request.args)
            return not page_name or not g.get("db_path") or get_page_visibility(get_db(), page_name)

        return {
            "csrf_token": session["_csrf_token"],
            "admin_page_access": admin_page_access,
            "page_is_visible": page_is_visible,
        }

    init_db(app.config["DB_PATH"], app.config)
    register_routes(app)
    return app


def get_db():
    if "db" not in g:
        g.db = get_conn(g.get("db_path"), g.get("db_config"))
    return g.db


def get_current_user():
    if "user_id" not in session:
        return AnonymousUser()
    if not hasattr(g, "_user"):
        row = get_db().execute("SELECT * FROM user WHERE id = ?", (session["user_id"],)).fetchone()
        g._user = CurrentUser(row) if row else AnonymousUser()
    return g._user


def _begin_db_transaction(db):
    if hasattr(db, "begin"):
        db.begin()
    else:
        db.execute("BEGIN IMMEDIATE")


def _rollback_db(db):
    if hasattr(db, "rollback"):
        db.rollback()
    else:
        db.execute("ROLLBACK")


def _session_device_identifier():
    address = request.remote_addr or "unknown"
    fingerprint = f"{request.user_agent.string}|{address}"
    return hashlib.sha256(fingerprint.encode("utf-8")).hexdigest()


def _client_device(db):
    address = request.remote_addr or "unknown"
    return db.execute(
        "SELECT id, status, is_active FROM device WHERE ip_address = ? ORDER BY id LIMIT 1",
        (address,),
    ).fetchone()


def _request_device_details():
    user_agent = request.user_agent.string.lower()
    if any(marker in user_agent for marker in ("ipad", "tablet", "kindle", "silk/")) or (
        "android" in user_agent and "mobile" not in user_agent
    ):
        device_type = "Tablet"
    elif any(marker in user_agent for marker in ("mobile", "iphone", "ipod", "windows phone")):
        device_type = "Mobile"
    else:
        device_type = "Desktop/Laptop"

    if "edg/" in user_agent:
        browser = "Edge"
    elif "opr/" in user_agent or "opera" in user_agent:
        browser = "Opera"
    elif "firefox/" in user_agent or "fxios/" in user_agent:
        browser = "Firefox"
    elif "chrome/" in user_agent or "crios/" in user_agent:
        browser = "Chrome"
    elif "safari/" in user_agent:
        browser = "Safari"
    else:
        browser = "Unknown"

    if "windows" in user_agent:
        operating_system = "Windows"
    elif "android" in user_agent:
        operating_system = "Android"
    elif "iphone" in user_agent or "ipad" in user_agent or "ipod" in user_agent:
        operating_system = "iOS/iPadOS"
    elif "mac os" in user_agent or "macintosh" in user_agent:
        operating_system = "macOS"
    elif "cros" in user_agent:
        operating_system = "ChromeOS"
    elif "linux" in user_agent:
        operating_system = "Linux"
    else:
        operating_system = "Unknown"

    return {
        "device_type": device_type,
        "browser": browser,
        "operating_system": operating_system,
        "ip_address": (request.remote_addr or "unknown")[:45],
    }


def _record_admin_login_activity(db, admin_id, status, session_id=None, login_time=None):
    details = _request_device_details()
    db.execute(
        "INSERT INTO login_activity "
        "(admin_id, device_type, browser, operating_system, ip_address, session_id, login_status, login_time) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?)",
        (
            admin_id,
            details["device_type"],
            details["browser"],
            details["operating_system"],
            details["ip_address"],
            session_id,
            status,
            login_time or utcnow_iso(),
        ),
    )


def _record_device_status_audit(
    db, device_id, student_id, previous_status, new_status, action,
    actor_id, actor_name, reason, action_time=None, ip_address=None,
):
    db.execute(
        "INSERT INTO device_status_audit "
        "(device_id, student_id, previous_status, new_status, action, admin_id, "
        "admin_name, reason, action_time, ip_address) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        (
            device_id,
            student_id,
            previous_status,
            new_status,
            action,
            actor_id,
            actor_name,
            reason,
            action_time or utcnow_iso(),
            (ip_address or request.remote_addr or "unknown")[:45],
        ),
    )


def _record_expired_device_unlocks(db, system_actor_id):
    current_time = utcnow_iso()
    expired_locks = db.execute(
        "SELECT id, current_student_id, lock_until FROM device "
        "WHERE lock_until IS NOT NULL AND lock_until <= ? AND current_student_id IS NOT NULL",
        (current_time,),
    ).fetchall()
    for device in expired_locks:
        updated = db.execute(
            "UPDATE device SET lock_until = NULL, current_student_id = NULL "
            "WHERE id = ? AND lock_until = ? AND current_student_id = ?",
            (device["id"], device["lock_until"], device["current_student_id"]),
        )
        if updated.rowcount:
            _record_device_status_audit(
                db,
                device["id"],
                device["current_student_id"],
                "LOCKED",
                "UNLOCKED",
                "UNLOCK",
                system_actor_id,
                "Automatic release",
                "The configured device lock duration elapsed.",
                action_time=device["lock_until"],
            )
    if expired_locks:
        db.commit()


def _claim_login_session(db, user_id, app):
    token = secrets.token_urlsafe(32)
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    now = utcnow_iso()
    expires_at = (
        datetime.now(timezone.utc).replace(tzinfo=None)
        + app.config["PERMANENT_SESSION_LIFETIME"]
    ).strftime("%Y-%m-%d %H:%M:%S")

    try:
        _begin_db_transaction(db)
        active = db.execute(
            "SELECT auth_session_id, token_hash, expires_at FROM active_user_session "
            "JOIN auth_session ON auth_session.id = active_user_session.auth_session_id "
            "WHERE active_user_session.user_id = ?",
            (user_id,),
        ).fetchone()
        user_row = db.execute("SELECT role FROM user WHERE id = ?", (user_id,)).fetchone()
        current_token = session.get("auth_token")
        current_token_hash = (
            hashlib.sha256(current_token.encode("utf-8")).hexdigest()
            if current_token
            else None
        )
        if active and current_token_hash == active["token_hash"]:
            _rollback_db(db)
            return current_token
        if active:
            db.execute(
                "UPDATE auth_session SET status = 'LOGGED_OUT', logout_time = ? WHERE id = ?",
                (now, active["auth_session_id"]),
            )
            if user_row and user_row["role"] == "admin":
                db.execute(
                    "UPDATE login_activity SET login_status = 'Replaced', logout_time = ? "
                    "WHERE session_id = ? AND logout_time IS NULL",
                    (now, active["token_hash"]),
                )
            db.execute("DELETE FROM active_user_session WHERE user_id = ?", (user_id,))

        result = db.execute(
            "INSERT INTO auth_session "
            "(user_id, token_hash, device_identifier, device_id, login_time, last_activity, status, expires_at) "
            "VALUES (?, ?, ?, ?, ?, ?, 'ACTIVE', ?)",
            (
                user_id,
                token_hash,
                _session_device_identifier(),
                _client_device(db)["id"] if user_row and user_row["role"] != "admin" and _client_device(db) else None,
                now,
                now,
                expires_at,
            ),
        )
        auth_session_id = result.lastrowid
        if user_row and user_row["role"] == "admin":
            _record_admin_login_activity(db, user_id, "Success", token_hash, now)
        db.execute(
            "INSERT INTO active_user_session (user_id, auth_session_id) VALUES (?, ?)",
            (user_id, auth_session_id),
        )
        db.commit()
        return token
    except Exception as error:
        _rollback_db(db)
        if is_integrity_error(error):
            return None
        raise


def _invalidate_auth_session(db, token, user_id=None):
    if not token:
        return
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    now = utcnow_iso()
    db.execute(
        "UPDATE auth_session SET status = 'LOGGED_OUT', logout_time = ? "
        "WHERE token_hash = ? AND status = 'ACTIVE'",
        (now, token_hash),
    )
    db.execute(
        "UPDATE login_activity SET login_status = 'Logged out', logout_time = ? "
        "WHERE session_id = ? AND logout_time IS NULL",
        (now, token_hash),
    )
    if user_id is None:
        db.execute(
            "DELETE FROM active_user_session WHERE auth_session_id IN "
            "(SELECT id FROM auth_session WHERE token_hash = ?)",
            (token_hash,),
        )
    else:
        db.execute("DELETE FROM active_user_session WHERE user_id = ?", (user_id,))
    db.commit()


def _has_valid_auth_session(db):
    token = session.get("auth_token")
    user_id = session.get("user_id")
    if not token or not user_id:
        return False
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    now = utcnow_iso()
    active = db.execute(
        "SELECT auth_session.id, auth_session.expires_at, auth_session.device_id, user.role, user.account_status, user.super_admin "
        "FROM auth_session JOIN user ON user.id = auth_session.user_id "
        "JOIN active_user_session ON active_user_session.auth_session_id = auth_session.id "
        "WHERE active_user_session.user_id = ? AND auth_session.token_hash = ? "
        "AND auth_session.status = 'ACTIVE'",
        (user_id, token_hash),
    ).fetchone()
    if not active:
        return False
    if active["expires_at"] < now:
        _invalidate_auth_session(db, token, user_id)
        return False
    if active["role"] == "admin" and (active["account_status"] or "active").lower() not in {"active"}:
        _invalidate_auth_session(db, token, user_id)
        flash("This admin account is inactive or suspended. Please contact the Super Admin.", "error")
        return False
    if active["role"] != "admin":
        device = None
        if active["device_id"]:
            device = db.execute(
                "SELECT id, status, ip_address FROM device WHERE id = ?",
                (active["device_id"],),
            ).fetchone()
        else:
            device = db.execute(
                "SELECT id, status, ip_address FROM device WHERE ip_address = ? ORDER BY id LIMIT 1",
                (request.remote_addr or "unknown",),
            ).fetchone()
        if device and (
            device["status"] != "APPROVED"
            or device["ip_address"] != (request.remote_addr or "unknown")
        ):
            session["_blocked_device_session"] = True
            _invalidate_auth_session(db, token, user_id)
            return False
    db.execute(
        "UPDATE auth_session SET last_activity = ? WHERE id = ?",
        (now, active["id"]),
    )
    db.commit()
    return True


def login_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        if not get_current_user().is_authenticated or not _has_valid_auth_session(get_db()):
            blocked_device = session.pop("_blocked_device_session", False)
            session.clear()
            if blocked_device:
                flash(
                    "This device has been blocked by the administrator. Please contact the administrator.",
                    "error",
                )
            return redirect(url_for("login"))
        return f(*args, **kwargs)
    return wrapper


def admin_required(f):
    @wraps(f)
    def wrapper(*args, **kwargs):
        user = get_current_user()
        if not user.is_authenticated or not user.is_admin:
            abort(403)
        page_name = page_name_for_admin_endpoint(request.endpoint, request.args)
        if page_name and not user.is_super_admin and not get_page_visibility(get_db(), page_name):
            abort(403)
        required_permission = required_admin_permission(request.endpoint, request.method)
        if required_permission and not admin_can_access_module(
            user, required_permission[0], required_permission[1], get_db()
        ):
            abort(403)
        return f(*args, **kwargs)
    return wrapper


def allowed_file(filename, app):
    if filename is None:
        return False
    filename = str(filename).strip()
    if not filename or "." not in filename:
        return False
    extension = filename.rsplit(".", 1)[1].lower()
    return extension in app.config.get("ALLOWED_EXTENSIONS", set())


def valid_course_id(db, raw_course_id):
    if raw_course_id is None:
        return None
    if isinstance(raw_course_id, str):
        raw_course_id = raw_course_id.strip()
    if raw_course_id in ("", None):
        return None
    try:
        course_id = int(raw_course_id)
    except (TypeError, ValueError):
        return None
    if course_id < 0 or not db.execute("SELECT 1 FROM course WHERE id = ?", (course_id,)).fetchone():
        return None
    return course_id


def valid_subject_id(db, raw_subject_id):
    if raw_subject_id is None:
        return None
    if isinstance(raw_subject_id, str):
        raw_subject_id = raw_subject_id.strip()
    if raw_subject_id in ("", None):
        return None
    try:
        subject_id = int(raw_subject_id)
    except (TypeError, ValueError):
        return None
    if subject_id < 0 or not db.execute("SELECT 1 FROM subject WHERE id = ?", (subject_id,)).fetchone():
        return None
    return subject_id


def valid_release_minutes(raw_minutes, default=15):
    try:
        minutes = int(raw_minutes)
    except (TypeError, ValueError):
        return default
    return max(1, min(minutes, 240))


def normalize_course_section(raw_value):
    value = (raw_value or "").strip()
    if not value:
        return ""
    section_map = {"1": "A", "2": "B", "3": "C", "4": "D", "5": "E", "6": "F", "7": "G"}
    upper_value = value.upper()
    if upper_value in {"A", "B", "C", "D", "E", "F", "G"}:
        return upper_value
    return section_map.get(value, "")


def default_admin_permissions():
    return {
        "dashboard": {"view": True, "add": False, "edit": False, "delete": False, "import": False, "export": False},
        "student_management": {"view": True, "add": True, "edit": True, "delete": True, "import": True, "export": True},
        "registration": {"view": True, "add": True, "edit": True, "delete": True, "import": True, "export": True},
        "course_management": {"view": True, "add": True, "edit": True, "delete": True, "import": True, "export": True},
        "subject_management": {"view": True, "add": True, "edit": True, "delete": True, "import": True, "export": True},
        "attendance": {"view": True, "add": True, "edit": True, "delete": True, "import": False, "export": True},
        "lab_management": {"view": True, "add": True, "edit": True, "delete": True, "import": False, "export": False},
        "lab_keys": {"view": True, "add": True, "edit": True, "delete": False, "import": False, "export": False},
        "documents": {"view": True, "add": True, "edit": True, "delete": True, "import": True, "export": True},
        "reports": {"view": True, "add": False, "edit": False, "delete": False, "import": False, "export": True},
        "user_management": {"view": True, "add": True, "edit": True, "delete": False, "import": False, "export": False},
        "admin_management": {"view": True, "add": True, "edit": True, "delete": True, "import": False, "export": False},
        "settings": {"view": True, "add": True, "edit": True, "delete": False, "import": False, "export": False},
    }


def normalize_admin_permissions(raw_permissions):
    defaults = default_admin_permissions()
    if not raw_permissions:
        return defaults
    if isinstance(raw_permissions, dict):
        data = raw_permissions
    else:
        try:
            data = json.loads(raw_permissions)
        except (TypeError, ValueError):
            return defaults
    if not isinstance(data, dict):
        return defaults
    for module_name, module_data in defaults.items():
        module_value = data.get(module_name, {})
        if not isinstance(module_value, dict):
            data[module_name] = dict(module_data)
            continue
        for permission_name in module_data:
            if permission_name not in module_value:
                module_value[permission_name] = bool(module_data[permission_name])
        data[module_name] = {name: bool(module_value.get(name, False)) for name in module_data}
    return data


def get_admin_permissions(db, admin_id):
    row = db.execute(
        "SELECT permissions FROM admin_permission WHERE admin_id = ?",
        (admin_id,),
    ).fetchone()
    if row and row["permissions"]:
        return normalize_admin_permissions(row["permissions"])
    return default_admin_permissions()


def save_admin_permissions(db, target_admin_id, permissions, updated_by_admin_id):
    payload = json.dumps(normalize_admin_permissions(permissions), separators=(",", ":"))
    existing = db.execute(
        "SELECT id FROM admin_permission WHERE admin_id = ?",
        (target_admin_id,),
    ).fetchone()
    now = utcnow_iso()
    if existing:
        db.execute(
            "UPDATE admin_permission SET permissions = ?, updated_by = ?, updated_at = ? WHERE admin_id = ?",
            (payload, updated_by_admin_id, now, target_admin_id),
        )
    else:
        db.execute(
            "INSERT INTO admin_permission (admin_id, permissions, updated_by, updated_at) VALUES (?, ?, ?, ?)",
            (target_admin_id, payload, updated_by_admin_id, now),
        )
    db.commit()
    return payload


def admin_can_access_module(user, module_name, action_name, db):
    if not user or not getattr(user, "is_authenticated", False):
        return False
    if getattr(user, "is_super_admin", False):
        return True
    if not getattr(user, "is_admin", False):
        return False
    permissions = get_admin_permissions(db, user.id)
    return bool(permissions.get(module_name, {}).get(action_name, False))


def required_admin_permission(endpoint, method):
    endpoint_permissions = {
        "admin_panel": ("dashboard", "view"),
        "admin_permissions_page": ("admin_management", "view"),
        "admin_login_activity": ("user_management", "view"),
        "admin_force_logout_session": ("user_management", "edit"),
        "admin_subjects": ("subject_management", "add" if method == "POST" else "view"),
        "admin_edit_subject": ("subject_management", "edit" if method == "POST" else "view"),
        "admin_remove_subject_course": ("subject_management", "delete"),
        "admin_delete_subject": ("subject_management", "delete"),
        "admin_assign_subject": ("subject_management", "edit"),
        "admin_devices": ("lab_management", "add" if method == "POST" else "view"),
        "admin_device_settings": ("lab_management", "edit"),
        "admin_toggle_device": ("lab_management", "edit"),
        "admin_update_device_timer": ("lab_management", "edit"),
        "admin_delete_device": ("lab_management", "delete"),
        "admin_release_device": ("lab_management", "edit"),
        "admin_students": ("student_management", "view"),
        "admin_logout_student_device": ("student_management", "edit"),
        "admin_import_students": ("student_management", "import"),
        "admin_export_students": ("student_management", "export"),
        "admin_edit_student": ("student_management", "edit" if method == "POST" else "view"),
        "admin_delete_student": ("student_management", "delete"),
        "admin_deactivate_student": ("student_management", "edit"),
        "admin_activate_student": ("student_management", "edit"),
        "admin_documents": ("documents", "view"),
        "admin_courses": ("course_management", "add" if method == "POST" else "view"),
        "admin_edit_course": ("course_management", "edit" if method == "POST" else "view"),
        "admin_delete_course": ("course_management", "delete"),
        "admin_upload_document": ("documents", "add"),
        "admin_show_document": ("documents", "view"),
        "admin_update_document": ("documents", "edit" if method == "POST" else "view"),
        "admin_delete_document": ("documents", "delete"),
        "admin_generate_lab_key": ("lab_keys", "add"),
        "admin_deactivate_lab": ("lab_keys", "edit"),
        "admin_attendance": ("attendance", "view"),
        "admin_reports": ("reports", "view"),
        "admin_export_attendance": ("attendance", "export"),
        "admin_view_attendance": ("attendance", "view"),
        "admin_password_page": ("settings", "edit"),
        "admin_forgot_user_password": ("user_management", "edit"),
    }
    return endpoint_permissions.get(endpoint)


def admin_permission_required(module_name, action_name):
    def decorator(func):
        @wraps(func)
        def wrapped(*args, **kwargs):
            user = get_current_user()
            db = get_db()
            if not user.is_authenticated or not user.is_admin:
                abort(403)
            if not admin_can_access_module(user, module_name, action_name, db):
                abort(403)
            return func(*args, **kwargs)
        return wrapped
    return decorator


def register_routes(app):

    @app.before_request
    def _stash_db_path():
        g.db_path = app.config["DB_PATH"]
        g.db_config = app.config

    @app.before_request
    def enforce_page_visibility():
        page_name = page_name_for_admin_endpoint(request.endpoint, request.args)
        if not page_name:
            return
        user = get_current_user()
        if getattr(user, "is_super_admin", False):
            return
        if not get_page_visibility(get_db(), page_name):
            abort(403)
        if user.is_admin:
            page_key = admin_page_key_for_endpoint(request.endpoint, request.args)
            if page_key and not get_admin_page_permissions(get_db(), user.id).get(page_key, False):
                abort(403)

    @app.teardown_appcontext
    def _close_db(exception=None):
        db = g.pop("db", None)
        if db is not None:
            db.close()

    @app.context_processor
    def inject_user():
        return {"current_user": get_current_user()}

    @app.route("/")
    def index():
        user = get_current_user()
        if user.is_authenticated:
            return redirect(url_for("admin_panel" if user.is_admin else "dashboard"))
        return redirect(url_for("login"))

    # ---------- Auth ----------

    @app.route("/register", methods=["GET", "POST"])
    def register():
        db = get_db()
        courses = db.execute(
            "SELECT * FROM course WHERE COALESCE(status, 'active') != 'inactive' ORDER BY name, year"
        ).fetchall()

        if request.method == "POST":
            name = request.form.get("name", "").strip()
            student_id = request.form.get("student_id", "").strip()
            section = request.form.get("section", "").strip()
            email = request.form.get("email", "").strip().lower()
            mobile_no = request.form.get("mobile_no", "")
            password = request.form.get("password", "")
            course_id = valid_course_id(db, request.form.get("course_id"))

            if not name or not email or not password or not student_id or not section or not mobile_no:
                flash("Student ID, name, section, email, mobile number, and password are required.", "error")
                return redirect(url_for("register"))

            if section not in {"A", "B", "C", "D", "E", "F", "G"}:
                flash("Select a valid section.", "error")
                return redirect(url_for("register"))

            if not is_valid_email(email):
                flash("Enter a valid email address.", "error")
                return redirect(url_for("register"))

            if not is_valid_mobile_number(mobile_no):
                flash("Mobile number must be exactly 10 digits (0-9).", "error")
                return redirect(url_for("register"))

            if len(password) < 6:
                flash("Password must be at least 6 characters.", "error")
                return redirect(url_for("register"))

            if course_id is None:
                flash("Select a valid course.", "error")
                return redirect(url_for("register"))

            course = db.execute(
                "SELECT year, status FROM course WHERE id = ?", (course_id,)
            ).fetchone()
            if not course or (course["status"] or "active").lower() == "inactive" or not course["year"]:
                flash("The selected course does not have a valid year.", "error")
                return redirect(url_for("register"))

            existing = db.execute("SELECT id FROM user WHERE email = ?", (email,)).fetchone()
            if existing:
                flash("An account with that email already exists.", "error")
                return redirect(url_for("register"))

            db.execute(
                "INSERT INTO user (name, student_id, section, year, email, mobile_no, password_hash, role, course_id, created_at) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, 'student', ?, ?)",
                (
                    name,
                    student_id,
                    section,
                    course["year"],
                    email,
                    mobile_no,
                    generate_password_hash(password),
                    course_id,
                    utcnow_iso(),
                ),
            )
            db.commit()
            flash("Registration successful. Please log in.", "success")
            return redirect(url_for("login"))

        return render_template("register.html", courses=courses)

    @app.route("/login", methods=["GET", "POST"])
    def login():
        db = get_db()
        courses = db.execute(
            "SELECT * FROM course WHERE COALESCE(status, 'active') != 'inactive' ORDER BY name"
        ).fetchall()

        if request.method == "POST":
            email = request.form.get("email", "").strip().lower()
            password = request.form.get("password", "")
            raw_course_id = request.form.get("course_id", "").strip()
            course_id = valid_course_id(db, raw_course_id)
            if not is_valid_email(email):
                flash("Enter a valid email address.", "error")
                return render_template(
                    "login.html",
                    courses=courses,
                    email=email,
                    selected_course_id=raw_course_id,
                )
            user_row = db.execute("SELECT * FROM user WHERE email = ?", (email,)).fetchone()

            if user_row and check_password_hash(user_row["password_hash"], password):
                if user_row["role"] == "admin" and (user_row["account_status"] or "active").lower() not in {"active"}:
                    flash("This admin account is inactive or suspended. Please contact the Super Admin.", "error")
                    return render_template(
                        "login.html",
                        courses=courses,
                        email=email,
                        selected_course_id=raw_course_id,
                    )
                is_active = user_row["is_active"] if "is_active" in user_row.keys() else 1
                if user_row["role"] != "admin" and is_active == 0:
                    flash("Your account is deactivated. Please contact the admin.", "error")
                    return render_template(
                        "login.html",
                        courses=courses,
                        email=email,
                        selected_course_id=raw_course_id,
                    )
                db.execute(
                    "UPDATE user SET failed_login_attempts = 0 WHERE id = ?",
                    (user_row["id"],),
                )
                db.commit()
                if user_row["role"] != "admin":
                    selected_course = db.execute(
                        "SELECT status FROM course WHERE id = ?", (course_id,)
                    ).fetchone() if course_id is not None else None
                    if raw_course_id and (course_id is None or course_id != user_row["course_id"]):
                        flash("Select your registered course.", "error")
                        return render_template(
                            "login.html",
                            courses=courses,
                            email=email,
                            selected_course_id=raw_course_id,
                        )
                    if selected_course and (selected_course["status"] or "active").lower() == "inactive":
                        flash("Select a valid course.", "error")
                        return render_template(
                            "login.html",
                            courses=courses,
                            email=email,
                            selected_course_id=raw_course_id,
                        )
                auth_token = _claim_login_session(db, user_row["id"], app)
                if not auth_token:
                    flash(
                        "This account is already logged in on another device. "
                        "Please logout from the previous device before logging in here.",
                        "error",
                    )
                    return render_template(
                        "login.html",
                        courses=courses,
                        email=email,
                        selected_course_id=raw_course_id,
                    )
                session.clear()
                session.permanent = True
                session["user_id"] = user_row["id"]
                session["auth_token"] = auth_token
                return redirect(url_for("admin_panel" if user_row["role"] == "admin" else "dashboard"))
            if user_row and user_row["role"] == "admin":
                _record_admin_login_activity(db, user_row["id"], "Failed")
                db.commit()
            if user_row and user_row["role"] != "admin":
                failed_attempts = (user_row["failed_login_attempts"] or 0) + 1
                if failed_attempts >= 4:
                    db.execute(
                        "UPDATE user SET failed_login_attempts = ?, is_active = 0 WHERE id = ?",
                        (failed_attempts, user_row["id"]),
                    )
                    db.commit()
                    flash("Your account has been deactivated after 4 incorrect password attempts. Please contact the admin.", "error")
                    return render_template(
                        "login.html",
                        courses=courses,
                        email=email,
                        selected_course_id=raw_course_id,
                    )
                db.execute(
                    "UPDATE user SET failed_login_attempts = ? WHERE id = ?",
                    (failed_attempts, user_row["id"]),
                )
                db.commit()
            flash("Invalid email or password.", "error")
        return render_template(
            "login.html",
            courses=courses,
            email=email if request.method == "POST" else "",
            selected_course_id=raw_course_id if request.method == "POST" else "",
        )

    @app.route("/logout", methods=["POST"])
    @login_required
    def logout():
        _invalidate_auth_session(get_db(), session.get("auth_token"), session.get("user_id"))
        session.clear()
        return redirect(url_for("login"))

    @app.route("/forgot-password", methods=["GET", "POST"])
    def forgot_password():
        if request.method == "POST":
            db = get_db()
            email = request.form.get("email", "").strip().lower()
            old_password = request.form.get("old_password", "")
            new_password = request.form.get("new_password", "")
            confirm_password = request.form.get("confirm_password", "")

            if not is_valid_email(email):
                flash("Enter a valid email address.", "error")
                return render_template("forgot_password.html")

            user_row = db.execute("SELECT * FROM user WHERE email = ?", (email,)).fetchone()
            if not user_row or user_row["role"] == "admin":
                flash("Student account not found.", "error")
                return render_template("forgot_password.html")

            if not check_password_hash(user_row["password_hash"], old_password):
                flash("Current password is incorrect.", "error")
                return render_template("forgot_password.html")

            if len(new_password) < 6:
                flash("New password must be at least 6 characters.", "error")
                return render_template("forgot_password.html")

            if new_password != confirm_password:
                flash("New password and confirmation do not match.", "error")
                return render_template("forgot_password.html")

            db.execute(
                "UPDATE user SET password_hash = ? WHERE id = ?",
                (generate_password_hash(new_password), user_row["id"]),
            )
            db.commit()
            flash("Password updated successfully. Please log in again.", "success")
            return redirect(url_for("login"))

        return render_template("forgot_password.html")

    def _process_password_form(target_override=None):
        target = request.form.get("target") or request.args.get("target") or target_override or "admin"

        if request.method == "POST":
            db = get_db()
            email = request.form.get("email", "").strip().lower()
            new_password = request.form.get("new_password", "")
            confirm_password = request.form.get("confirm_password", "")

            user_row = (
                db.execute("SELECT * FROM user WHERE email = ?", (email,)).fetchone()
                if is_valid_email(email)
                else None
            )
            if not is_valid_email(email):
                flash("Enter a valid email address.", "error")
            elif not user_row:
                flash("Account not found.", "error")
            elif user_row["super_admin"]:
                flash("The super admin password cannot be reset.", "error")
            elif target == "admin" and user_row["role"] != "admin":
                flash("Admin account not found.", "error")
            elif target == "student" and user_row["role"] == "admin":
                flash("Student account not found.", "error")
            elif len(new_password) < 6:
                flash("New password must be at least 6 characters.", "error")
            elif new_password != confirm_password:
                flash("New password and confirmation do not match.", "error")
            else:
                db.execute(
                    "UPDATE user SET password_hash = ? WHERE id = ?",
                    (generate_password_hash(new_password), user_row["id"]),
                )
                db.commit()
                if target == "admin":
                    flash("Admin password updated successfully. Please log in again.", "success")
                    return redirect(url_for("login"))
                flash("User password updated successfully.", "success")
                return redirect(url_for("admin_panel"))

        return render_template("admin_password.html")

    @app.route("/admin/permissions", methods=["GET", "POST"])
    @login_required
    @admin_required
    def admin_permissions_page():
        db = get_db()
        current_user = get_current_user()
        if not current_user.is_super_admin:
            abort(403)

        edit_admin_id = request.args.get("edit_id", "").strip()
        view_admin_id = request.args.get("view_id", "").strip()
        selected_status = request.args.get("status", "").strip().lower()
        selected_role = request.args.get("role", "").strip().lower()
        search_term = request.args.get("search", "").strip()

        if request.method == "POST":
            action = request.form.get("action")
            if action == "create_admin":
                name = request.form.get("name", "").strip()
                email = request.form.get("email", "").strip().lower()
                password = request.form.get("password", "")
                if not name or len(name) > 255:
                    flash("Enter an admin name under 256 characters.", "error")
                elif not is_valid_email(email):
                    flash("Enter a valid admin email address.", "error")
                elif len(password) < 6:
                    flash("Admin passwords must be at least 6 characters.", "error")
                elif db.execute("SELECT 1 FROM user WHERE email = ?", (email,)).fetchone():
                    flash("An account with that email already exists.", "error")
                else:
                    db.execute(
                        "INSERT INTO user (name, email, password_hash, role, super_admin, course_id, is_active, account_status, created_at) "
                        "VALUES (?, ?, ?, 'admin', 0, NULL, 1, 'active', ?)",
                        (name, email, generate_password_hash(password), utcnow_iso()),
                    )
                    target_id = db.execute("SELECT id FROM user WHERE email = ?", (email,)).fetchone()["id"]
                    db.execute(
                        "INSERT INTO admin_permission_audit (super_admin_id, target_admin_id, module_name, permission_name, old_value, new_value, changed_at) "
                        "VALUES (?, ?, 'admin_management', 'account_created', 0, 1, ?)",
                        (current_user.id, target_id, utcnow_iso()),
                    )
                    db.commit()
                    flash("Admin account created.", "success")
                    return redirect(url_for("admin_permissions_page"))
            elif action == "save_permissions":
                target_admin_id = request.form.get("admin_id")
                if target_admin_id and target_admin_id.isdigit():
                    target_admin_id = int(target_admin_id)
                    target = db.execute(
                        "SELECT id, super_admin, email, name FROM user WHERE id = ? AND role = 'admin'",
                        (target_admin_id,),
                    ).fetchone()
                    if target and not target["super_admin"]:
                        raw_permissions = {}
                        for module_name, module_data in default_admin_permissions().items():
                            raw_permissions[module_name] = {}
                            for perm_name in module_data:
                                raw_permissions[module_name][perm_name] = request.form.get(f"perm_{module_name}_{perm_name}") == "on"
                        old_permissions = get_admin_permissions(db, target_admin_id)
                        save_admin_permissions(db, target_admin_id, raw_permissions, current_user.id)
                        for module_name, module_data in default_admin_permissions().items():
                            for perm_name in module_data:
                                old_value = bool(old_permissions.get(module_name, {}).get(perm_name, False))
                                new_value = bool(raw_permissions.get(module_name, {}).get(perm_name, False))
                                if old_value != new_value:
                                    db.execute(
                                        "INSERT INTO admin_permission_audit (super_admin_id, target_admin_id, module_name, permission_name, old_value, new_value, changed_at) VALUES (?, ?, ?, ?, ?, ?, ?)",
                                        (current_user.id, target_admin_id, module_name, perm_name, 1 if old_value else 0, 1 if new_value else 0, utcnow_iso()),
                                    )
                        db.commit()
                        flash("Admin permissions saved.", "success")
                        return redirect(url_for("admin_permissions_page"))
                    flash("Super admin permissions cannot be edited from this page.", "error")
                    return redirect(url_for("admin_permissions_page"))
            elif action == "set_status":
                target_admin_id = request.form.get("admin_id")
                new_status = (request.form.get("account_status") or "active").strip().lower()
                if target_admin_id and target_admin_id.isdigit() and new_status in {"active", "inactive", "suspended"}:
                    target = db.execute(
                        "SELECT id, super_admin, email, account_status FROM user WHERE id = ? AND role = 'admin'",
                        (int(target_admin_id),),
                    ).fetchone()
                    if target and not target["super_admin"]:
                        old_status = (target["account_status"] or "active").lower()
                        if old_status != new_status:
                            status_values = {"inactive": 0, "active": 1, "suspended": 2}
                            now = utcnow_iso()
                            db.execute(
                                "UPDATE user SET account_status = ? WHERE id = ?",
                                (new_status, int(target_admin_id)),
                            )
                            db.execute(
                                "INSERT INTO admin_permission_audit (super_admin_id, target_admin_id, module_name, permission_name, old_value, new_value, changed_at) "
                                "VALUES (?, ?, 'admin_management', 'account_status', ?, ?, ?)",
                                (current_user.id, int(target_admin_id), status_values.get(old_status, 0), status_values[new_status], now),
                            )
                            if new_status != "active":
                                db.execute(
                                    "UPDATE auth_session SET status = 'FORCE_LOGOUT', logout_time = ? "
                                    "WHERE user_id = ? AND status = 'ACTIVE'",
                                    (now, int(target_admin_id)),
                                )
                                db.execute("DELETE FROM active_user_session WHERE user_id = ?", (int(target_admin_id),))
                                db.execute(
                                    "UPDATE login_activity SET login_status = 'Force logged out', logout_time = ? "
                                    "WHERE admin_id = ? AND logout_time IS NULL",
                                    (now, int(target_admin_id)),
                                )
                        db.commit()
                        flash(f"Admin status updated to {new_status}.", "success")
                        return redirect(url_for("admin_permissions_page"))
                    flash("Only non-super-admin accounts can be updated here.", "error")
                    return redirect(url_for("admin_permissions_page"))
            elif action == "delete_admin":
                target_admin_id = request.form.get("admin_id")
                if target_admin_id and target_admin_id.isdigit():
                    target_id = int(target_admin_id)
                    target = db.execute(
                        "SELECT id, super_admin, email FROM user WHERE id = ? AND role = 'admin'",
                        (target_id,),
                    ).fetchone()
                    if target and not target["super_admin"] and target_id != current_user.id:
                        db.execute("DELETE FROM admin_permission WHERE admin_id = ?", (target_id,))
                        db.execute("DELETE FROM auth_session WHERE user_id = ?", (target_id,))
                        db.execute("DELETE FROM active_user_session WHERE user_id = ?", (target_id,))
                        db.execute("DELETE FROM user WHERE id = ?", (target_id,))
                        db.commit()
                        flash("Admin account deleted.", "success")
                        return redirect(url_for("admin_permissions_page"))
                    flash("That admin cannot be deleted from here.", "error")
                    return redirect(url_for("admin_permissions_page"))
            elif action == "update_page_visibility":
                page_id = request.form.get("page_id", "").strip()
                visible_value = request.form.get("is_visible", "").strip()
                if page_id.isdigit() and visible_value in {"0", "1"}:
                    page = db.execute(
                        "SELECT id FROM page_visibility WHERE id = ?", (int(page_id),)
                    ).fetchone()
                    if page:
                        db.execute(
                            "UPDATE page_visibility SET is_visible = ?, updated_by = ?, updated_at = ? WHERE id = ?",
                            (int(visible_value), current_user.id, utcnow_iso(), int(page_id)),
                        )
                        db.commit()
                        flash("Page visibility updated successfully.", "success")
                        return redirect(url_for(
                            "admin_permissions_page",
                            visibility_search=request.form.get("visibility_search", ""),
                            visibility_filter=request.form.get("visibility_filter", ""),
                            _anchor="page-visibility",
                        ))
                flash("Select a valid page visibility setting.", "error")
                return redirect(url_for("admin_permissions_page", _anchor="page-visibility"))
            elif action == "save_admin_page_permissions":
                target_admin_id = request.form.get("admin_id", "").strip()
                if target_admin_id.isdigit():
                    target_admin_id = int(target_admin_id)
                    target = db.execute(
                        "SELECT id, super_admin FROM user WHERE id = ? AND role = 'admin'",
                        (target_admin_id,),
                    ).fetchone()
                    if target and not target["super_admin"]:
                        old_permissions = get_admin_page_permissions(db, target_admin_id)
                        for page_key, _page_name, _page_url, _module_name, _action_name in ADMIN_PAGE_PERMISSION_DEFINITIONS:
                            is_enabled = request.form.get(f"page_perm_{page_key}") == "on"
                            existing = db.execute(
                                "SELECT id FROM admin_page_permission WHERE admin_id = ? AND page_key = ?",
                                (target_admin_id, page_key),
                            ).fetchone()
                            if existing:
                                db.execute(
                                    "UPDATE admin_page_permission SET is_enabled = ?, updated_by = ?, updated_at = ? WHERE id = ?",
                                    (int(is_enabled), current_user.id, utcnow_iso(), existing["id"]),
                                )
                            else:
                                db.execute(
                                    "INSERT INTO admin_page_permission (admin_id, page_key, is_enabled, updated_by, updated_at) VALUES (?, ?, ?, ?, ?)",
                                    (target_admin_id, page_key, int(is_enabled), current_user.id, utcnow_iso()),
                                )
                            if old_permissions.get(page_key, False) != is_enabled:
                                db.execute(
                                    "INSERT INTO admin_permission_audit (super_admin_id, target_admin_id, module_name, permission_name, old_value, new_value, changed_at) VALUES (?, ?, 'page_access', ?, ?, ?, ?)",
                                    (current_user.id, target_admin_id, page_key, int(old_permissions.get(page_key, False)), int(is_enabled), utcnow_iso()),
                                )
                        db.commit()
                        flash("Page permissions saved.", "success")
                        return redirect(url_for(
                            "admin_permissions_page",
                            pages_id=target_admin_id,
                            search=request.form.get("search", ""),
                            role=request.form.get("role", ""),
                            status=request.form.get("status", ""),
                            _anchor="page-permissions",
                        ))
                flash("Select a valid admin account.", "error")
                return redirect(url_for("admin_permissions_page"))

        admins = db.execute(
            "SELECT u.*, COALESCE((SELECT login_time FROM login_activity WHERE admin_id = u.id ORDER BY login_time DESC LIMIT 1), '') AS last_login "
            "FROM user u WHERE u.role = 'admin' ORDER BY u.name"
        ).fetchall()
        admin_stats = {
            "total": len(admins),
            "super_admins": sum(1 for admin in admins if admin["super_admin"]),
            "active": sum(1 for admin in admins if (admin["account_status"] or "active").lower() == "active"),
            "inactive": sum(1 for admin in admins if (admin["account_status"] or "active").lower() == "inactive"),
            "suspended": sum(1 for admin in admins if (admin["account_status"] or "active").lower() == "suspended"),
        }

        filtered_admins = []
        for admin in admins:
            if search_term and search_term.lower() not in (admin["name"] or "").lower() and search_term.lower() not in (admin["email"] or "").lower():
                continue
            if selected_status and (admin["account_status"] or "active").lower() != selected_status:
                continue
            if selected_role and (
                (admin["super_admin"] == 1 and selected_role != "super_admin")
                or (admin["super_admin"] == 0 and selected_role != "admin")
            ):
                continue
            permissions = get_admin_permissions(db, admin["id"])
            active_permission_count = sum(
                1
                for module_permissions in permissions.values()
                for enabled in module_permissions.values()
                if enabled
            )
            filtered_admins.append({
                "admin": admin,
                "permissions": permissions,
                "active_permission_count": active_permission_count,
            })

        permission_modal = None
        modal_id = edit_admin_id or view_admin_id or next(
            (entry["admin"]["id"] for entry in filtered_admins if not entry["admin"]["super_admin"]),
            None,
        )
        if modal_id:
            modal_admin = db.execute(
                "SELECT * FROM user WHERE id = ? AND role = 'admin'",
                (int(modal_id),),
            ).fetchone()
            if modal_admin:
                permission_modal = {
                    "admin": modal_admin,
                    "permissions": get_admin_permissions(db, modal_admin["id"]),
                    "view_only": bool(view_admin_id and not edit_admin_id),
                }

        page_permissions_admin = None
        page_permission_rows = []
        selected_pages_admin_id = request.args.get("pages_id", "").strip()
        if selected_pages_admin_id.isdigit():
            selected_pages_admin = db.execute(
                "SELECT id, name, email, super_admin FROM user WHERE id = ? AND role = 'admin'",
                (int(selected_pages_admin_id),),
            ).fetchone()
            if selected_pages_admin and not selected_pages_admin["super_admin"]:
                page_permissions_admin = selected_pages_admin
                saved_page_permissions = get_admin_page_permissions(db, selected_pages_admin["id"])
                page_permission_rows = [
                    {
                        "key": page_key,
                        "name": page_name,
                        "url": page_url,
                        "is_enabled": saved_page_permissions[page_key],
                    }
                    for page_key, page_name, page_url, _module_name, _action_name
                    in ADMIN_PAGE_PERMISSION_DEFINITIONS
                ]

        visibility_search = request.args.get("visibility_search", "").strip()
        visibility_filter = request.args.get("visibility_filter", "").strip().lower()
        page_visibility = db.execute(
            "SELECT * FROM page_visibility ORDER BY id"
        ).fetchall()
        if visibility_search:
            page_visibility = [
                page for page in page_visibility
                if visibility_search.lower() in page["page_name"].lower()
            ]
        if visibility_filter in {"visible", "hidden"}:
            expected_visibility = 1 if visibility_filter == "visible" else 0
            page_visibility = [
                page for page in page_visibility
                if bool(page["is_visible"]) == bool(expected_visibility)
            ]

        return render_template(
            "admin_permissions.html",
            admins=filtered_admins,
            permission_modal=permission_modal,
            selected_status=selected_status,
            selected_role=selected_role,
            search_term=search_term,
            default_permissions=default_admin_permissions(),
            admin_stats=admin_stats,
            page_visibility=page_visibility,
            visibility_search=visibility_search,
            visibility_filter=visibility_filter,
            page_permissions_admin=page_permissions_admin,
            page_permission_rows=page_permission_rows,
        )

    @app.route("/admin/password", methods=["GET", "POST"])
    @login_required
    @admin_required
    def admin_password_page():
        return _process_password_form()

    @app.route("/admin/forgot-password", methods=["GET", "POST"])
    def admin_forgot_password():
        return _process_password_form("admin")

    @app.route("/admin/forgot-user-password", methods=["GET", "POST"])
    @login_required
    @admin_required
    def admin_forgot_user_password():
        return _process_password_form("student")

    # ---------- Student ----------

    @app.route("/dashboard")
    @login_required
    def dashboard():
        user = get_current_user()
        if user.is_admin:
            return redirect(url_for("admin_panel"))

        db = get_db()
        course = None
        sessions = []
        documents = []
        selected_month = request.args.get("month", datetime.now(
            timezone(timedelta(hours=5, minutes=30))
        ).strftime("%Y-%m"))
        try:
            month_date = datetime.strptime(selected_month, "%Y-%m")
        except ValueError:
            selected_month = datetime.now(timezone(timedelta(hours=5, minutes=30))).strftime("%Y-%m")
            month_date = datetime.strptime(selected_month, "%Y-%m")
        days_in_month = calendar.monthrange(month_date.year, month_date.month)[1]
        days_in_month = calendar.monthrange(month_date.year, month_date.month)[1]
        attendance_marks = {}
        if user.course_id:
            course = db.execute("SELECT * FROM course WHERE id = ?", (user.course_id,)).fetchone()
            sessions = db.execute(
                "SELECT * FROM lab_session WHERE course_id = ? AND is_active = 1 AND expires_at >= ? "
                "ORDER BY created_at DESC",
                (user.course_id, utcnow_iso()),
            ).fetchall()
            documents = db.execute(
                "SELECT document.*, course.name AS course_name, course.year AS course_year "
                "FROM document JOIN course ON document.course_id = course.id "
                "WHERE document.course_id = ? ORDER BY document.uploaded_at DESC",
                (user.course_id,),
            ).fetchall()
            attendance_rows = db.execute(
                "SELECT attendance.timestamp, lab_session.code AS lab_code "
                "FROM attendance JOIN lab_session ON attendance.lab_session_id = lab_session.id "
                "WHERE attendance.student_id = ? AND lab_session.course_id = ?",
                (user.id, user.course_id),
            ).fetchall()
            for attendance in attendance_rows:
                try:
                    utc_value = datetime.strptime(
                        attendance["timestamp"][:19], "%Y-%m-%d %H:%M:%S"
                    ).replace(tzinfo=timezone.utc)
                    ist_value = utc_value.astimezone(timezone(timedelta(hours=5, minutes=30)))
                except (TypeError, ValueError):
                    continue
                if ist_value.strftime("%Y-%m") == selected_month:
                    attendance_marks[ist_value.day] = {
                        "time": ist_value.strftime("%H:%M"),
                        "key": attendance["lab_code"],
                    }
        subject_values = [value.strip() for value in (course["subject"] if course else "").split("|") if value.strip()]
        present_count = len(attendance_marks)
        return render_template(
            "dashboard.html",
            course=course,
            sessions=sessions,
            documents=documents,
            subject_values=subject_values,
            selected_month=selected_month,
            days=range(1, days_in_month + 1),
            attendance_marks=attendance_marks,
            present_count=present_count,
            absent_count=days_in_month - present_count,
            attendance_percentage=round((present_count / days_in_month) * 100) if days_in_month else 0,
        )

    @app.route("/attendance")
    @login_required
    def attendance():
        user = get_current_user()
        if user.is_admin:
            return redirect(url_for("admin_attendance"))

        db = get_db()
        course = None
        attendance_marks = {}
        subject_attendance = []
        selected_month = request.args.get(
            "month", datetime.now(timezone(timedelta(hours=5, minutes=30))).strftime("%Y-%m")
        )
        try:
            month_date = datetime.strptime(selected_month, "%Y-%m")
        except ValueError:
            selected_month = datetime.now(timezone(timedelta(hours=5, minutes=30))).strftime("%Y-%m")
            month_date = datetime.strptime(selected_month, "%Y-%m")
        days_in_month = calendar.monthrange(month_date.year, month_date.month)[1]
        available_months = [
            (f"{month_date.year}-{month:02d}", datetime(month_date.year, month, 1).strftime("%B %Y"))
            for month in range(1, 13)
        ]

        if user.course_id:
            course = db.execute("SELECT * FROM course WHERE id = ?", (user.course_id,)).fetchone()
            subject_rows = db.execute(
                "SELECT subject.id, subject.name FROM subject "
                "JOIN course_subject ON course_subject.subject_id = subject.id "
                "WHERE course_subject.course_id = ? ORDER BY subject.name",
                (user.course_id,),
            ).fetchall()
            attendance_rows = db.execute(
                "SELECT attendance.timestamp, lab_session.code AS lab_code, "
                "lab_session.subject_id, subject.name AS subject_name "
                "FROM attendance JOIN lab_session ON attendance.lab_session_id = lab_session.id "
                "LEFT JOIN subject ON subject.id = lab_session.subject_id "
                "WHERE attendance.student_id = ? AND lab_session.course_id = ?",
                (user.id, user.course_id),
            ).fetchall()
            subject_data = {
                subject_row["id"]: {"name": subject_row["name"], "marks": {}}
                for subject_row in subject_rows
            }
            for attendance_row in attendance_rows:
                try:
                    utc_value = datetime.strptime(
                        attendance_row["timestamp"][:19], "%Y-%m-%d %H:%M:%S"
                    ).replace(tzinfo=timezone.utc)
                    ist_value = utc_value.astimezone(timezone(timedelta(hours=5, minutes=30)))
                except (TypeError, ValueError):
                    continue
                if ist_value.strftime("%Y-%m") == selected_month:
                    mark = {
                        "time": ist_value.strftime("%H:%M"),
                        "key": attendance_row["lab_code"],
                    }
                    attendance_marks[ist_value.day] = mark
                    subject_id = attendance_row["subject_id"] or 0
                    if subject_id not in subject_data:
                        subject_data[subject_id] = {
                            "name": attendance_row["subject_name"] or "General Attendance",
                            "marks": {},
                        }
                    subject_data[subject_id]["marks"][ist_value.day] = mark

            for subject_data_row in subject_data.values():
                present = len(subject_data_row["marks"])
                subject_attendance.append(
                    {
                        "name": subject_data_row["name"],
                        "marks": subject_data_row["marks"],
                        "present_count": present,
                        "absent_count": days_in_month - present,
                        "percentage": round((present / days_in_month) * 100) if days_in_month else 0,
                    }
                )

        present_count = len(attendance_marks)
        subject_present_count = sum(subject_record["present_count"] for subject_record in subject_attendance)
        subject_absent_count = sum(subject_record["absent_count"] for subject_record in subject_attendance)
        subject_attendance_percentage = round(
            (subject_present_count / (subject_present_count + subject_absent_count)) * 100
        ) if subject_present_count + subject_absent_count else 0
        subject = course["subject"] if course and course["subject"] else "-"
        return render_template(
            "attendance.html",
            course=course,
            subject=subject,
            selected_month=selected_month,
            available_months=available_months,
            month_label=month_date.strftime("%B %Y"),
            days=range(1, days_in_month + 1),
            attendance_marks=attendance_marks,
            subject_attendance=subject_attendance,
            subject_present_count=subject_present_count,
            subject_absent_count=subject_absent_count,
            subject_attendance_percentage=subject_attendance_percentage,
            present_count=present_count,
            absent_count=days_in_month - present_count,
            attendance_percentage=round((present_count / days_in_month) * 100) if days_in_month else 0,
        )

    @app.route("/documents")
    @login_required
    def documents():
        user = get_current_user()
        db = get_db()
        if user.is_admin:
            docs = db.execute(
                "SELECT document.*, course.name AS course_name, course.year AS course_year FROM document "
                "JOIN course ON document.course_id = course.id ORDER BY document.uploaded_at DESC"
            ).fetchall()
        elif user.course_id:
            docs = db.execute(
                "SELECT document.*, course.name AS course_name, course.year AS course_year FROM document "
                "JOIN course ON document.course_id = course.id "
                "WHERE document.course_id = ? ORDER BY document.uploaded_at DESC",
                (user.course_id,),
            ).fetchall()
        else:
            docs = []
        return render_template("documents.html", documents=docs)

    @app.route("/documents/download/<int:doc_id>")
    @login_required
    def download_document(doc_id):
        user = get_current_user()
        db = get_db()
        doc = db.execute("SELECT * FROM document WHERE id = ?", (doc_id,)).fetchone()
        if not doc:
            abort(404)
        if not user.is_admin and doc["course_id"] != user.course_id:
            abort(403)
        return send_from_directory(app.config["UPLOAD_FOLDER"], doc["filename"], as_attachment=True)

    @app.route("/lab/enter", methods=["GET", "POST"])
    @login_required
    def lab_enter():
        user = get_current_user()
        if user.is_admin:
            abort(403)

        if request.method == "POST":
            db = get_db()
            code = request.form.get("code", "").strip()

            session_row = db.execute(
                "SELECT * FROM lab_session WHERE code = ? AND course_id = ? "
                "ORDER BY created_at DESC LIMIT 1",
                (code, user.course_id),
            ).fetchone()

            is_valid = (
                session_row is not None
                and session_row["is_active"] == 1
                and session_row["created_at"] <= utcnow_iso()
                and session_row["expires_at"] >= utcnow_iso()
            )
            if not is_valid:
                flash("Invalid or expired lab key.", "error")
                return redirect(url_for("lab_enter"))

            device_ip = request.remote_addr or "unknown"
            current_time = utcnow_iso()
            device = db.execute(
                "SELECT id, is_active, lock_until, release_minutes FROM device WHERE ip_address = ? ORDER BY id LIMIT 1",
                (device_ip,),
            ).fetchone()
            if not device:
                flash("Attendance is only available from a registered lab PC.", "error")
                return redirect(url_for("lab_enter"))
            if not device["is_active"]:
                flash("This PC has been deactivated and cannot mark attendance.", "error")
                return redirect(url_for("lab_enter"))
            if device and device["lock_until"] and device["lock_until"] > current_time:
                flash(
                    f"This PC is locked until {device['lock_until']} because attendance was already marked.",
                    "error",
                )
                return redirect(url_for("lab_enter"))
            existing = db.execute(
                "SELECT id FROM attendance WHERE student_id = ? AND lab_session_id = ?",
                (user.id, session_row["id"]),
            ).fetchone()
            if existing:
                flash("Attendance already recorded for this session.", "error")
                return redirect(url_for("lab_enter"))
            lock_until = (
                datetime.strptime(current_time, "%Y-%m-%d %H:%M:%S")
                + timedelta(minutes=device["release_minutes"] or app.config.get("DEVICE_LOCK_MINUTES", 15))
            ).strftime("%Y-%m-%d %H:%M:%S")
            db.execute(
                "UPDATE device SET lock_until = ?, current_student_id = ? WHERE id = ?",
                (lock_until, user.id, device["id"]),
            )
            _record_device_status_audit(
                db,
                device["id"],
                user.id,
                "UNLOCKED",
                "LOCKED",
                "LOCK",
                user.id,
                user.name,
                "Attendance recorded; the device is locked until the release timer expires.",
                action_time=current_time,
                ip_address=device_ip,
            )

            db.execute(
                "INSERT INTO attendance (student_id, lab_session_id, timestamp) VALUES (?, ?, ?)",
                (user.id, session_row["id"], utcnow_iso()),
            )
            db.commit()
            flash("Attendance marked successfully.", "success")
            return redirect(url_for("dashboard"))

        return redirect(url_for("dashboard"))

    # ---------- Admin ----------

    @app.route("/admin")
    @login_required
    @admin_required
    def admin_panel():
        db = get_db()
        courses = db.execute("SELECT * FROM course ORDER BY name").fetchall()
        subjects = db.execute("SELECT * FROM subject ORDER BY name").fetchall()
        students_count = db.execute("SELECT COUNT(*) AS c FROM user WHERE role = 'student'").fetchone()["c"]
        docs_count = db.execute("SELECT COUNT(*) AS c FROM document").fetchone()["c"]
        documents = db.execute(
            "SELECT document.*, course.name AS course_name, course.year AS course_year "
            "FROM document JOIN course ON document.course_id = course.id "
            "ORDER BY document.uploaded_at DESC"
        ).fetchall()
        active_sessions_raw = db.execute(
            "SELECT lab_session.*, course.name AS course_name, course.year AS course_year, "
            "course.subject AS course_subject, subject.name AS subject_name "
            "FROM lab_session "
            "JOIN course ON lab_session.course_id = course.id "
            "LEFT JOIN subject ON lab_session.subject_id = subject.id "
            "WHERE lab_session.is_active = 1 AND lab_session.expires_at >= ? "
            "ORDER BY lab_session.created_at DESC",
            (utcnow_iso(),),
        ).fetchall()
        active_sessions = []
        for session_row in active_sessions_raw:
            session_dict = dict(session_row)
            session_dict["validity_minutes"] = max(
                0,
                int(
                    (
                        datetime.strptime(str(session_dict["expires_at"]).replace("T", " ")[:19], "%Y-%m-%d %H:%M:%S")
                        - datetime.strptime(str(session_dict["created_at"]).replace("T", " ")[:19], "%Y-%m-%d %H:%M:%S")
                    ).total_seconds() // 60
                ),
            )
            active_sessions.append(session_dict)
        admin_sessions = db.execute(
            "SELECT auth_session.id AS session_id, auth_session.last_activity, auth_session.login_time, "
            "login_activity.device_type, login_activity.browser, login_activity.ip_address, "
            "COALESCE(NULLIF(user.student_id, ''), user.email) AS admin_identifier "
            "FROM auth_session "
            "JOIN active_user_session ON active_user_session.auth_session_id = auth_session.id "
            "JOIN user ON user.id = auth_session.user_id "
            "JOIN login_activity ON login_activity.session_id = auth_session.token_hash "
            "WHERE user.role = 'admin' AND auth_session.status = 'ACTIVE' AND auth_session.expires_at >= ? "
            "AND login_activity.logout_time IS NULL ORDER BY auth_session.last_activity DESC",
            (utcnow_iso(),),
        ).fetchall()
        return render_template(
            "admin.html",
            courses=courses,
            students_count=students_count,
            docs_count=docs_count,
            documents=documents,
            active_sessions=active_sessions,
            admin_sessions=admin_sessions,
            subjects=subjects,
        )

    @app.route("/admin/login-activity")
    @login_required
    @admin_required
    def admin_login_activity():
        db = get_db()
        _record_expired_device_unlocks(db, get_current_user().id)
        activity = db.execute(
            "SELECT auth_session.login_time, auth_session.logout_time, auth_session.status AS session_status, "
            "user.name AS user_name, user.student_id, course.name AS course_name, course.year AS course_year "
            "FROM auth_session JOIN user ON user.id = auth_session.user_id "
            "LEFT JOIN course ON course.id = user.course_id "
            "WHERE user.role = 'student' "
            "ORDER BY auth_session.login_time DESC, auth_session.id DESC LIMIT 500"
        ).fetchall()
        device_activity = db.execute(
            "SELECT device_status_audit.*, device.name AS device_name, device.asset_tag, "
            "user.name AS student_name, user.student_id AS student_identifier, "
            "user.email AS student_email, "
            "user.mobile_no AS student_mobile "
            "FROM device_status_audit "
            "LEFT JOIN device ON device.id = device_status_audit.device_id "
            "LEFT JOIN user ON user.id = device_status_audit.student_id "
            "WHERE device_status_audit.action IN ('LOCK', 'UNLOCK', 'DELETE') "
            "ORDER BY device_status_audit.action_time DESC, device_status_audit.id DESC LIMIT 500"
        ).fetchall()
        return render_template(
            "admin_login_activity.html",
            activity=activity,
            device_activity=device_activity,
        )

    @app.route("/admin/sessions/<int:session_id>/logout", methods=["POST"])
    @login_required
    @admin_required
    def admin_force_logout_session(session_id):
        db = get_db()
        active = db.execute(
            "SELECT auth_session.id, auth_session.token_hash FROM auth_session "
            "JOIN active_user_session ON active_user_session.auth_session_id = auth_session.id "
            "JOIN user ON user.id = auth_session.user_id "
            "WHERE auth_session.id = ? AND user.role = 'admin' AND auth_session.status = 'ACTIVE'",
            (session_id,),
        ).fetchone()
        if not active:
            abort(404)
        now = utcnow_iso()
        db.execute(
            "UPDATE auth_session SET status = 'FORCE_LOGOUT', logout_time = ? WHERE id = ?",
            (now, session_id),
        )
        db.execute("DELETE FROM active_user_session WHERE auth_session_id = ?", (session_id,))
        db.execute(
            "UPDATE login_activity SET login_status = 'Force logged out', logout_time = ? "
            "WHERE session_id = ? AND logout_time IS NULL",
            (now, active["token_hash"]),
        )
        db.commit()
        flash("Admin session logged out.", "success")
        return redirect(url_for("admin_panel"))

    @app.route("/admin/subjects", methods=["GET", "POST"])
    @login_required
    @admin_required
    def admin_subjects():
        db = get_db()
        if request.method == "POST":
            name = request.form.get("name", "").strip()
            subject_id = request.form.get("subject_id", "").strip()
            if subject_id.isdigit():
                duplicate = db.execute(
                    "SELECT id FROM subject WHERE LOWER(name) = LOWER(?) AND id != ?",
                    (name, int(subject_id)),
                ).fetchone()
                if not name:
                    flash("Subject name is required.", "error")
                elif duplicate:
                    flash("That subject already exists.", "error")
                else:
                    db.execute("UPDATE subject SET name = ? WHERE id = ?", (name, int(subject_id)))
                    db.commit()
                    flash("Subject updated.", "success")
                return redirect(url_for("admin_subjects"))
            if not name:
                flash("Subject name is required.", "error")
            elif db.execute("SELECT 1 FROM subject WHERE LOWER(name) = LOWER(?)", (name,)).fetchone():
                flash("That subject already exists.", "error")
            else:
                db.execute("INSERT INTO subject (name) VALUES (?)", (name,))
                db.commit()
                flash("Subject added.", "success")
            return redirect(url_for("admin_subjects"))

        subject_rows = db.execute(
            "SELECT subject.*, COUNT(DISTINCT course_subject.course_id) AS course_count "
            "FROM subject LEFT JOIN course_subject ON course_subject.subject_id = subject.id "
            "GROUP BY subject.id ORDER BY subject.name"
        ).fetchall()
        subjects = []
        for subject_row in subject_rows:
            subject = dict(subject_row)
            subject["assigned_courses"] = db.execute(
                "SELECT course.id, course.name, course.year FROM course "
                "JOIN course_subject ON course_subject.course_id = course.id "
                "WHERE course_subject.subject_id = ? ORDER BY course.name, course.year",
                (subject["id"],),
            ).fetchall()
            subjects.append(subject)
        courses = db.execute("SELECT * FROM course ORDER BY name").fetchall()
        return render_template("admin_subjects.html", subjects=subjects, courses=courses)

    @app.route("/admin/subjects/<int:subject_id>/edit", methods=["GET", "POST"])
    @login_required
    @admin_required
    def admin_edit_subject(subject_id):
        db = get_db()
        subject = db.execute("SELECT * FROM subject WHERE id = ?", (subject_id,)).fetchone()
        if not subject:
            abort(404)
        assigned_courses = db.execute(
            "SELECT course.id, course.name, course.year FROM course "
            "JOIN course_subject ON course_subject.course_id = course.id "
            "WHERE course_subject.subject_id = ? ORDER BY course.name, course.year",
            (subject_id,),
        ).fetchall()
        if request.method == "POST":
            name = request.form.get("name", "").strip()
            duplicate = db.execute(
                "SELECT id FROM subject WHERE LOWER(name) = LOWER(?) AND id != ?",
                (name, subject_id),
            ).fetchone()
            if not name:
                flash("Subject name is required.", "error")
            elif duplicate:
                flash("That subject already exists.", "error")
            else:
                db.execute("UPDATE subject SET name = ? WHERE id = ?", (name, subject_id))
                db.commit()
                flash("Subject updated.", "success")
                return redirect(url_for("admin_subjects"))
        return render_template(
            "admin_subject_edit.html",
            subject=subject,
            assigned_courses=assigned_courses,
        )

    @app.route("/admin/subjects/<int:subject_id>/courses/<int:course_id>/delete", methods=["POST"])
    @login_required
    @admin_required
    def admin_remove_subject_course(subject_id, course_id):
        db = get_db()
        if not db.execute("SELECT id FROM subject WHERE id = ?", (subject_id,)).fetchone():
            abort(404)
        db.execute(
            "DELETE FROM course_subject WHERE subject_id = ? AND course_id = ?",
            (subject_id, course_id),
        )
        db.commit()
        flash("Subject removed from course.", "success")
        return redirect(url_for("admin_subjects"))

    @app.route("/admin/subjects/<int:subject_id>/delete", methods=["POST"])
    @login_required
    @admin_required
    def admin_delete_subject(subject_id):
        db = get_db()
        subject = db.execute("SELECT id FROM subject WHERE id = ?", (subject_id,)).fetchone()
        if not subject:
            abort(404)
        db.execute("DELETE FROM subject WHERE id = ?", (subject_id,))
        db.commit()
        flash("Subject deleted.", "success")
        return redirect(url_for("admin_subjects"))

    @app.route("/admin/subjects/<int:subject_id>/courses", methods=["POST"])
    @login_required
    @admin_required
    def admin_assign_subject(subject_id):
        db = get_db()
        if not db.execute("SELECT 1 FROM subject WHERE id = ?", (subject_id,)).fetchone():
            abort(404)
        course_id = valid_course_id(db, request.form.get("course_id"))
        if course_id is None:
            flash("Select a valid course.", "error")
        else:
            if not db.execute(
                "SELECT 1 FROM course_subject WHERE course_id = ? AND subject_id = ?",
                (course_id, subject_id),
            ).fetchone():
                db.execute(
                    "INSERT INTO course_subject (course_id, subject_id) VALUES (?, ?)",
                    (course_id, subject_id),
                )
            db.commit()
            flash("Subject assigned to course.", "success")
        return redirect(url_for("admin_subjects"))

    @app.route("/admin/devices", methods=["GET", "POST"])
    @login_required
    @admin_required
    def admin_devices():
        db = get_db()
        if request.method == "POST":
            name = request.form.get("name", "").strip()
            asset_tag = request.form.get("asset_tag", "").strip()
            ip_address = request.form.get("ip_address", "").strip()
            location = request.form.get("location", "").strip()
            release_minutes = valid_release_minutes(
                request.form.get("release_minutes"), app.config.get("DEVICE_LOCK_MINUTES", 15)
            )

            if not name or not asset_tag:
                flash("Device name and asset tag are required.", "error")
            elif db.execute(
                "SELECT 1 FROM device WHERE LOWER(asset_tag) = LOWER(?)", (asset_tag,)
            ).fetchone():
                flash("That asset tag is already in use.", "error")
            else:
                db.execute(
                    "INSERT INTO device (name, asset_tag, ip_address, location, release_minutes, created_at) "
                    "VALUES (?, ?, ?, ?, ?, ?)",
                    (
                        name,
                        asset_tag,
                        ip_address or None,
                        location or socket.gethostname(),
                        release_minutes,
                        utcnow_iso(),
                    ),
                )
                db.commit()
                flash("Device added.", "success")
            return redirect(url_for("admin_devices"))

        _record_expired_device_unlocks(db, get_current_user().id)
        devices = db.execute(
            "SELECT device.*, "
            "CASE WHEN device.asset_tag LIKE 'AUTO-PC-%' "
            "THEN 'PC_' || substr(device.asset_tag, 9) ELSE device.name END AS display_name, "
            "user.name AS current_student_name "
            "FROM device LEFT JOIN user ON user.id = device.current_student_id "
            "ORDER BY device.name, device.asset_tag"
        ).fetchall()
        current_time = utcnow_iso()
        active_locks = [
            device["lock_until"] for device in devices
            if device["lock_until"] and device["lock_until"] > current_time
        ]
        device_stats = {
            "total": len(devices),
            "active": sum(1 for device in devices if device["is_active"]),
            "deactivated": sum(1 for device in devices if not device["is_active"]),
            "locked": sum(
                1 for device in devices
                if device["lock_until"] and device["lock_until"] > current_time
            ),
        }
        return render_template(
            "admin_devices.html",
            devices=devices,
            device_stats=device_stats,
            lock_minutes=app.config.get("DEVICE_LOCK_MINUTES", 15),
            next_lock_until=min(active_locks) if active_locks else None,
            now_utc=utcnow_iso(),
        )

    @app.route("/admin/devices/settings", methods=["POST"])
    @login_required
    @admin_required
    def admin_device_settings():
        minutes = valid_release_minutes(request.form.get("lock_minutes"), app.config.get("DEVICE_LOCK_MINUTES", 15))
        app.config["DEVICE_LOCK_MINUTES"] = minutes
        db = get_db()
        db.execute("UPDATE device SET release_minutes = ?", (minutes,))
        db.commit()
        flash(f"PC lock duration saved: {minutes} minutes.", "success")
        return redirect(url_for("admin_devices"))

    @app.route("/admin/devices/<int:device_id>/toggle", methods=["POST"])
    @login_required
    @admin_required
    def admin_toggle_device(device_id):
        db = get_db()
        device = db.execute("SELECT id, is_active FROM device WHERE id = ?", (device_id,)).fetchone()
        if not device:
            abort(404)
        db.execute(
            "UPDATE device SET is_active = ? WHERE id = ?",
            (0 if device["is_active"] else 1, device_id),
        )
        db.commit()
        flash("Device status updated.", "success")
        return redirect(url_for("admin_devices"))

    @app.route("/admin/devices/<int:device_id>/timer", methods=["POST"])
    @login_required
    @admin_required
    def admin_update_device_timer(device_id):
        db = get_db()
        if not db.execute("SELECT id FROM device WHERE id = ?", (device_id,)).fetchone():
            abort(404)
        release_minutes = valid_release_minutes(
            request.form.get("release_minutes"), app.config.get("DEVICE_LOCK_MINUTES", 15)
        )
        db.execute(
            "UPDATE device SET release_minutes = ? WHERE id = ?",
            (release_minutes, device_id),
        )
        db.commit()
        flash("Device auto-release timer updated.", "success")
        return redirect(url_for("admin_devices"))

    @app.route("/admin/devices/<int:device_id>/delete", methods=["POST"])
    @login_required
    @admin_required
    def admin_delete_device(device_id):
        db = get_db()
        device = db.execute(
            "SELECT id, name, asset_tag, is_active, lock_until, current_student_id "
            "FROM device WHERE id = ?",
            (device_id,),
        ).fetchone()
        if not device:
            abort(404)
        admin = get_current_user()
        now = utcnow_iso()
        previous_status = (
            "LOCKED"
            if device["lock_until"] and device["lock_until"] > now
            else "ACTIVE" if device["is_active"] else "DEACTIVATED"
        )
        _record_device_status_audit(
            db,
            device_id,
            device["current_student_id"],
            previous_status,
            "DELETED",
            "DELETE",
            admin.id,
            admin.name,
            f"Device deleted: {device['name']} (asset tag {device['asset_tag']}).",
            action_time=now,
        )
        db.execute("DELETE FROM device WHERE id = ?", (device_id,))
        db.commit()
        flash("Device deleted.", "success")
        return redirect(url_for("admin_devices"))

    @app.route("/admin/devices/<int:device_id>/release", methods=["POST"])
    @login_required
    @admin_required
    def admin_release_device(device_id):
        db = get_db()
        device = db.execute(
            "SELECT id, lock_until, current_student_id FROM device WHERE id = ?",
            (device_id,),
        ).fetchone()
        if not device:
            abort(404)
        if device["lock_until"]:
            admin = get_current_user()
            _record_device_status_audit(
                db,
                device_id,
                device["current_student_id"],
                "LOCKED",
                "UNLOCKED",
                "UNLOCK",
                admin.id,
                admin.name,
                "Device lock released by administrator.",
            )
        db.execute(
            "UPDATE device SET lock_until = NULL, current_student_id = NULL WHERE id = ?",
            (device_id,),
        )
        db.commit()
        flash("Device lock released by admin.", "success")
        return redirect(url_for("admin_devices"))

    @app.route("/admin/students")
    @login_required
    @admin_required
    def admin_students():
        db = get_db()
        courses = db.execute("SELECT * FROM course ORDER BY name").fetchall()
        available_years = [
            row["year"]
            for row in db.execute(
                "SELECT DISTINCT substr(created_at, 1, 4) AS year "
                "FROM user WHERE role = 'student' ORDER BY year DESC"
            ).fetchall()
        ]
        selected_course_id = request.args.get("course_id", "").strip()
        search_term = request.args.get("search", "").strip()
        selected_year = request.args.get("year", "").strip()
        course_id = valid_course_id(db, selected_course_id)
        student_query = (
            "SELECT user.*, course.name AS course_name, course.year AS course_year, "
            "auth_session.id AS active_session_id, auth_session.login_time AS active_login_time "
            "FROM user LEFT JOIN course ON user.course_id = course.id "
            "LEFT JOIN active_user_session ON active_user_session.user_id = user.id "
            "LEFT JOIN auth_session ON auth_session.id = active_user_session.auth_session_id "
            "WHERE user.role = 'student'"
        )
        student_params = []
        if selected_course_id:
            student_query += " AND user.course_id = ?"
            student_params.append(course_id if course_id is not None else -1)
        if search_term:
            student_query += (
                " AND (user.student_id LIKE ? OR user.email LIKE ? OR user.mobile_no LIKE ?)"
            )
            search_pattern = f"%{search_term}%"
            student_params.extend([search_pattern, search_pattern, search_pattern])
        if selected_year.isdigit() and len(selected_year) == 4:
            student_query += " AND user.created_at LIKE ?"
            student_params.append(f"{selected_year}-%")
        student_query += " ORDER BY user.name"
        students = db.execute(
            student_query, student_params
        ).fetchall()
        return render_template(
            "admin_students.html",
            students=students,
            courses=courses,
            available_years=available_years,
            selected_course_id=selected_course_id,
            search_term=search_term,
            selected_year=selected_year,
        )

    @app.route("/admin/students/<int:student_id>/logout-device", methods=["POST"])
    @login_required
    @admin_required
    def admin_logout_student_device(student_id):
        db = get_db()
        student = db.execute(
            "SELECT id FROM user WHERE id = ? AND role = 'student'", (student_id,)
        ).fetchone()
        if not student:
            abort(404)
        now = utcnow_iso()
        db.execute(
            "UPDATE auth_session SET status = 'LOGGED_OUT', logout_time = ? "
            "WHERE id = (SELECT auth_session_id FROM active_user_session WHERE user_id = ?)",
            (now, student_id),
        )
        db.execute("DELETE FROM active_user_session WHERE user_id = ?", (student_id,))
        db.commit()
        flash("Student device session logged out.", "success")
        return redirect(url_for("admin_students"))

    @app.route("/admin/students/import", methods=["POST"])
    @login_required
    @admin_required
    def admin_import_students():
        uploaded_file = request.files.get("student_file")
        if not uploaded_file or not uploaded_file.filename:
            flash("Choose an Excel or CSV file to import.", "error")
            return redirect(url_for("admin_students"))

        try:
            rows = read_student_import(uploaded_file)
        except (ValueError, UnicodeDecodeError, xlrd.XLRDError, openpyxl.utils.exceptions.InvalidFileException) as exc:
            flash(f"Import failed: {exc}", "error")
            return redirect(url_for("admin_students"))

        if not rows:
            flash("Import failed: the file is empty.", "error")
            return redirect(url_for("admin_students"))

        header_map = {
            normalize_import_header(value): index
            for index, value in enumerate(rows[0])
            if value is not None and str(value).strip()
        }
        missing_columns = [
            column for column in IMPORT_COLUMNS
            if normalize_import_header(column) not in header_map
        ]
        if missing_columns:
            flash(f"Import failed. Missing required columns: {', '.join(missing_columns)}.", "error")
            return redirect(url_for("admin_students"))

        db = get_db()
        courses = db.execute("SELECT id, name, year FROM course ORDER BY name, year").fetchall()
        course_lookup = build_course_lookup(courses)
        course_name_ids = {}
        for course in courses:
            course_name = str(course["name"]).strip().lower()
            course_name_ids.setdefault(course_name, set()).add(course["id"])
        for course_name, course_ids in course_name_ids.items():
            if len(course_ids) == 1:
                course_lookup.setdefault(course_name, []).append(next(iter(course_ids)))

        imported_count = 0
        duplicate_ids = []
        errors = []
        seen_ids = set()

        def cell_text(row, column, strip=True):
            index = header_map[normalize_import_header(column)]
            if index >= len(row) or row[index] is None:
                return ""
            value = row[index]
            if isinstance(value, float) and value.is_integer():
                return str(int(value))
            text = str(value)
            return text.strip() if strip else text

        for row_number, row in enumerate(rows[1:], start=2):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            student_id = cell_text(row, "Student ID")
            name = cell_text(row, "Student Name")
            course_value = cell_text(row, "Course")
            mobile = cell_text(row, "Mobile", strip=False)
            email = cell_text(row, "Email").lower()
            if not all((student_id, name, course_value, mobile, email)):
                errors.append(f"Row {row_number}: all required values are needed.")
                continue
            if not is_valid_email(email):
                errors.append(f"Row {row_number}: email {email} is invalid.")
                continue
            if not is_valid_mobile_number(mobile):
                errors.append(f"Row {row_number}: mobile number {mobile} must be exactly 10 digits (0-9).")
                continue
            if student_id.lower() in seen_ids:
                duplicate_ids.append(student_id)
                continue
            seen_ids.add(student_id.lower())
            if db.execute(
                "SELECT id FROM user WHERE LOWER(student_id) = LOWER(?) AND role = 'student'",
                (student_id,),
            ).fetchone():
                duplicate_ids.append(student_id)
                continue
            if db.execute(
                "SELECT id FROM user WHERE LOWER(email) = LOWER(?)", (email,)
            ).fetchone():
                errors.append(f"Row {row_number}: email {email} already exists.")
                continue
            course_id = resolve_course_id(course_lookup, course_value)
            if course_id is None:
                errors.append(f"Row {row_number}: course {course_value} was not found.")
                continue
            section = cell_text(row, "Section").upper() if "section" in header_map else ""
            if section and section not in {"A", "B", "C", "D", "E", "F", "G"}:
                errors.append(f"Row {row_number}: section {section} is invalid.")
                continue
            course_year = db.execute(
                "SELECT year FROM course WHERE id = ?", (course_id,)
            ).fetchone()["year"]
            db.execute(
                "INSERT INTO user (name, student_id, section, year, email, mobile_no, password_hash, role, course_id, created_at) "
                "VALUES (?, ?, ?, ?, ?, ?, ?, 'student', ?, ?)",
                (
                    name,
                    student_id,
                    section,
                    course_year,
                    email,
                    mobile,
                    generate_password_hash(student_id),
                    course_id,
                    utcnow_iso(),
                ),
            )
            imported_count += 1

        db.commit()
        result_parts = [f"{imported_count} student(s) imported."]
        if duplicate_ids:
            result_parts.append(f"Duplicate Student IDs skipped: {', '.join(duplicate_ids[:10])}.")
        if errors:
            result_parts.append(f"{len(errors)} row(s) rejected: {' '.join(errors[:3])}")
        result_parts.append("Imported students use their Student ID as the initial password.")
        flash(" ".join(result_parts), "success" if imported_count else "error")
        return redirect(url_for("admin_students"))

    @app.route("/admin/students/export")
    @login_required
    @admin_required
    def admin_export_students():
        db = get_db()
        students = db.execute(
            "SELECT user.student_id, user.name, user.section, course.name AS course_name, course.year AS course_year, user.mobile_no, "
            "user.email, user.created_at, user.is_active "
            "FROM user LEFT JOIN course ON user.course_id = course.id "
            "WHERE user.role = 'student' ORDER BY user.name"
        ).fetchall()
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Students"
        headers = [
            "Student ID", "Student Name", "Course", "Section", "Mobile", "Email",
            "Registration Date", "Status",
        ]
        worksheet.append(headers)
        for student in students:
            worksheet.append([
                student["student_id"] or "",
                student["name"],
                format_course_label(student["course_name"], student["course_year"]),
                student["section"] or "",
                student["mobile_no"] or "",
                student["email"],
                student["created_at"] or "",
                "Active" if student["is_active"] else "Inactive",
            ])
        for cell in worksheet[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        widths = [16, 24, 24, 12, 16, 32, 24, 12]
        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name="students.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.route("/admin/students/<int:student_id>/edit", methods=["GET", "POST"])
    @login_required
    @admin_required
    def admin_edit_student(student_id):
        db = get_db()
        student = db.execute(
            "SELECT * FROM user WHERE id = ? AND role = 'student'", (student_id,)
        ).fetchone()
        if not student:
            abort(404)

        courses = db.execute("SELECT * FROM course ORDER BY name").fetchall()
        if request.method == "POST":
            name = request.form.get("name", "").strip()
            student_code = request.form.get("student_id", "").strip()
            email = request.form.get("email", "").strip().lower()
            mobile_no = request.form.get("mobile_no", "")
            section = request.form.get("section", "").strip().upper()
            course_id = valid_course_id(db, request.form.get("course_id"))
            course = db.execute("SELECT year FROM course WHERE id = ?", (course_id,)).fetchone() if course_id is not None else None

            if not is_valid_email(email):
                flash("Enter a valid email address.", "error")
                return render_template("admin_student_edit.html", student=student, courses=courses)

            if not is_valid_mobile_number(mobile_no):
                flash("Mobile number must be exactly 10 digits (0-9).", "error")
                return render_template("admin_student_edit.html", student=student, courses=courses)

            duplicate = db.execute(
                "SELECT id FROM user WHERE email = ? AND id != ?", (email, student_id)
            ).fetchone()
            if (not name or not student_code or not email or not mobile_no
                    or section not in {"A", "B", "C", "D", "E", "F", "G"}
                    or course is None or not course["year"]):
                flash("All student fields, a valid section, and a valid course are required.", "error")
            elif duplicate:
                flash("An account with that email already exists.", "error")
            else:
                db.execute(
                    "UPDATE user SET name = ?, student_id = ?, email = ?, mobile_no = ?, section = ?, year = ?, course_id = ? "
                    "WHERE id = ?",
                    (name, student_code, email, mobile_no, section, course["year"], course_id, student_id),
                )
                db.commit()
                flash("Student updated.", "success")
                return redirect(url_for("admin_students"))

        return render_template("admin_student_edit.html", student=student, courses=courses)

    @app.route("/admin/students/<int:student_id>/delete", methods=["POST"])
    @login_required
    @admin_required
    def admin_delete_student(student_id):
        db = get_db()
        student = db.execute(
            "SELECT id FROM user WHERE id = ? AND role = 'student'", (student_id,)
        ).fetchone()
        if not student:
            abort(404)

        db.execute("DELETE FROM attendance WHERE student_id = ?", (student_id,))
        db.execute("DELETE FROM user WHERE id = ?", (student_id,))
        db.commit()
        flash("Student deleted.", "success")
        return redirect(url_for("admin_students"))

    @app.route("/admin/students/<int:student_id>/deactivate", methods=["POST"])
    @login_required
    @admin_required
    def admin_deactivate_student(student_id):
        db = get_db()
        student = db.execute(
            "SELECT id FROM user WHERE id = ? AND role = 'student'", (student_id,)
        ).fetchone()
        if not student:
            abort(404)

        db.execute("UPDATE user SET is_active = 0 WHERE id = ?", (student_id,))
        db.commit()
        flash("Student deactivated.", "success")
        return redirect(url_for("admin_students"))

    @app.route("/admin/students/<int:student_id>/activate", methods=["POST"])
    @login_required
    @admin_required
    def admin_activate_student(student_id):
        db = get_db()
        student = db.execute(
            "SELECT id FROM user WHERE id = ? AND role = 'student'", (student_id,)
        ).fetchone()
        if not student:
            abort(404)

        db.execute("UPDATE user SET is_active = 1 WHERE id = ?", (student_id,))
        db.commit()
        flash("Student activated.", "success")
        return redirect(url_for("admin_students"))

    @app.route("/admin/documents")
    @login_required
    @admin_required
    def admin_documents():
        db = get_db()
        courses = db.execute("SELECT * FROM course ORDER BY name").fetchall()
        documents = db.execute(
            "SELECT document.*, course.name AS course_name, course.year AS course_year "
            "FROM document JOIN course ON document.course_id = course.id "
            "ORDER BY document.uploaded_at DESC"
        ).fetchall()
        return render_template("admin_documents.html", documents=documents, courses=courses)

    @app.route("/admin/courses", methods=["GET", "POST"])
    @login_required
    @admin_required
    def admin_courses():
        db = get_db()
        if request.method == "GET":
            search = request.args.get("q", "").strip()
            subject = request.args.get("subject", "").strip()
            query = "SELECT course.* FROM course"
            params = []
            filters = []
            if search:
                search_value = f"%{search}%"
                filters.append(
                    "(course.name LIKE ? OR course.subject LIKE ? OR course.year LIKE ? "
                    "OR EXISTS (SELECT 1 FROM course_subject cs JOIN subject s ON s.id = cs.subject_id "
                    "WHERE cs.course_id = course.id AND s.name LIKE ?))"
                )
                params.extend([search_value, search_value, search_value, search_value])
            if subject:
                subject_value = f"%{subject}%"
                filters.append(
                    "(LOWER(COALESCE(course.subject, '')) LIKE LOWER(?) OR EXISTS (SELECT 1 FROM course_subject cs "
                    "JOIN subject s ON s.id = cs.subject_id WHERE cs.course_id = course.id AND LOWER(s.name) LIKE LOWER(?)))"
                )
                params.extend([subject_value, subject_value])
            if filters:
                query += " WHERE " + " AND ".join(filters)
            query += " ORDER BY course.name"
            courses = db.execute(query, params).fetchall()
            subjects = db.execute("SELECT name FROM subject ORDER BY name").fetchall()
            return render_template(
                "admin_courses.html",
                courses=courses,
                search=search,
                subject=subject,
                subjects=subjects,
            )

        name = request.form.get("name", "").strip()
        year = request.form.get("year", "").strip()
        status = request.form.get("status", "").strip().lower()
        subject = request.form.get("subject", "").strip()
        course_section = normalize_course_section(request.form.get("course_section", ""))
        duplicate = db.execute(
            "SELECT id FROM course WHERE LOWER(TRIM(name)) = LOWER(?) AND year = ?",
            (name, year),
        ).fetchone() if name and year else None

        if (name and year in {"1Year", "2Year", "3Year", "4Year"}
                and course_section
                and status in {"active", "inactive"}
                and not duplicate):
            now = utcnow_iso()
            db.execute(
                "INSERT INTO course (name, subject, year, course_section, status, created_at, updated_at) "
                "VALUES (?, ?, ?, ?, ?, ?, ?)",
                (name, subject, year, course_section, status, now, now),
            )
            db.commit()
            flash("Course added.", "success")
        elif duplicate:
            flash("A course with this name and year already exists.", "error")
        else:
            flash("Please enter Course Name, Year, Section and Status.", "error")
        return redirect(url_for("admin_courses"))

    @app.route("/admin/courses/<int:course_id>/edit", methods=["GET", "POST"])
    @login_required
    @admin_required
    def admin_edit_course(course_id):
        db = get_db()
        course = db.execute("SELECT * FROM course WHERE id = ?", (course_id,)).fetchone()
        if not course:
            abort(404)

        if request.method == "POST":
            name = request.form.get("name", "").strip()
            year = request.form.get("year", "").strip()
            status = request.form.get("status", "").strip().lower()
            subject = request.form.get("subject", course["subject"] or "").strip()
            course_section = normalize_course_section(request.form.get("course_section", ""))
            duplicate = db.execute(
                "SELECT id FROM course WHERE LOWER(TRIM(name)) = LOWER(?) AND year = ? AND id != ?",
                (name, year, course_id),
            ).fetchone() if name and year else None
            if (not name or year not in {"1Year", "2Year", "3Year", "4Year"}
                    or not course_section
                    or status not in {"active", "inactive"}):
                flash("Please enter Course Name, Year, Section and Status.", "error")
            elif duplicate:
                flash("A course with this name and year already exists.", "error")
            else:
                db.execute(
                    "UPDATE course SET name = ?, subject = ?, year = ?, course_section = ?, "
                    "status = ?, updated_at = ? WHERE id = ?",
                    (name, subject, year, course_section, status, utcnow_iso(), course_id),
                )
                db.commit()
                flash("Course updated.", "success")
                return redirect(url_for("admin_courses"))

        return render_template("admin_course_edit.html", course=course)

    @app.route("/admin/courses/<int:course_id>/delete", methods=["POST"])
    @login_required
    @admin_required
    def admin_delete_course(course_id):
        db = get_db()
        course = db.execute("SELECT id FROM course WHERE id = ?", (course_id,)).fetchone()
        if not course:
            abort(404)

        documents = db.execute(
            "SELECT filename FROM document WHERE course_id = ?", (course_id,)
        ).fetchall()
        session_rows = db.execute(
            "SELECT id FROM lab_session WHERE course_id = ?", (course_id,)
        ).fetchall()

        for document in documents:
            file_path = os.path.join(app.config["UPLOAD_FOLDER"], document["filename"])
            if os.path.exists(file_path):
                os.remove(file_path)

        for session_row in session_rows:
            db.execute(
                "DELETE FROM attendance WHERE lab_session_id = ?", (session_row["id"],)
            )
        db.execute("DELETE FROM lab_session WHERE course_id = ?", (course_id,))
        db.execute("DELETE FROM document WHERE course_id = ?", (course_id,))
        db.execute("UPDATE user SET course_id = NULL WHERE course_id = ?", (course_id,))
        db.execute("DELETE FROM course WHERE id = ?", (course_id,))
        db.commit()
        flash("Course deleted. Related documents and lab sessions were also removed.", "success")
        return redirect(url_for("admin_courses"))

    @app.route("/admin/documents/upload", methods=["POST"])
    @login_required
    @admin_required
    def admin_upload_document():
        db = get_db()
        title = request.form.get("title", "").strip()
        course_id = valid_course_id(db, request.form.get("course_id"))
        file = request.files.get("file")

        if not (title and course_id and file and file.filename and allowed_file(file.filename, app)):
            flash("Missing title/course, or unsupported file type.", "error")
            return redirect(url_for("admin_panel"))

        filename = secure_filename(f"{secrets.token_hex(4)}_{file.filename}")
        file.save(os.path.join(app.config["UPLOAD_FOLDER"], filename))
        db.execute(
            "INSERT INTO document (course_id, title, filename, uploaded_at) VALUES (?, ?, ?, ?)",
            (course_id, title, filename, utcnow_iso()),
        )
        db.commit()
        flash("Document uploaded.", "success")
        return redirect(url_for("admin_documents"))

    @app.route("/admin/documents/<int:doc_id>/show")
    @login_required
    @admin_required
    def admin_show_document(doc_id):
        db = get_db()
        doc = db.execute(
            "SELECT document.*, course.name AS course_name, course.year AS course_year "
            "FROM document JOIN course ON document.course_id = course.id WHERE document.id = ?",
            (doc_id,),
        ).fetchone()
        if not doc:
            abort(404)
        return render_template("admin_document_show.html", document=doc)

    @app.route("/admin/documents/<int:doc_id>/edit", methods=["GET", "POST"])
    @app.route("/admin/documents/<int:doc_id>/update", methods=["GET", "POST"])
    @login_required
    @admin_required
    def admin_update_document(doc_id):
        db = get_db()
        doc = db.execute(
            "SELECT document.*, course.name AS course_name, course.year AS course_year "
            "FROM document JOIN course ON document.course_id = course.id WHERE document.id = ?",
            (doc_id,),
        ).fetchone()
        if not doc:
            abort(404)

        courses = db.execute("SELECT * FROM course ORDER BY name").fetchall()
        if request.method == "POST":
            title = request.form.get("title", "").strip()
            course_id = valid_course_id(db, request.form.get("course_id"))
            file = request.files.get("file")

            if not title or course_id is None:
                flash("Title and a valid course are required.", "error")
                return redirect(url_for("admin_update_document", doc_id=doc_id))

            update_values = [course_id, title]
            update_sql = "UPDATE document SET course_id = ?, title = ?"

            if file and file.filename:
                if not allowed_file(file.filename, app):
                    flash("Unsupported file type.", "error")
                    return redirect(url_for("admin_update_document", doc_id=doc_id))
                new_filename = secure_filename(f"{secrets.token_hex(4)}_{file.filename}")
                file.save(os.path.join(app.config["UPLOAD_FOLDER"], new_filename))
                old_path = os.path.join(app.config["UPLOAD_FOLDER"], doc["filename"])
                if os.path.exists(old_path):
                    os.remove(old_path)
                update_values.extend([new_filename, utcnow_iso(), doc_id])
                update_sql += " , filename = ?, uploaded_at = ? WHERE id = ?"
            else:
                update_values.append(doc_id)
                update_sql += " WHERE id = ?"

            db.execute(update_sql, update_values)
            db.commit()
            flash("Document updated.", "success")
            return redirect(url_for("admin_show_document", doc_id=doc_id))

        return render_template("admin_document_edit.html", document=doc, courses=courses)

    @app.route("/admin/documents/<int:doc_id>/delete", methods=["POST"])
    @login_required
    @admin_required
    def admin_delete_document(doc_id):
        db = get_db()
        doc = db.execute("SELECT * FROM document WHERE id = ?", (doc_id,)).fetchone()
        if not doc:
            abort(404)

        file_path = os.path.join(app.config["UPLOAD_FOLDER"], doc["filename"])
        if os.path.exists(file_path):
            os.remove(file_path)

        db.execute("DELETE FROM document WHERE id = ?", (doc_id,))
        db.commit()
        flash("Document deleted.", "success")
        return redirect(url_for("admin_panel"))

    @app.route("/admin/lab/generate", methods=["POST"])
    @login_required
    @admin_required
    def admin_generate_lab_key():
        db = get_db()
        course_id = valid_course_id(db, request.form.get("course_id"))
        subject_id = valid_subject_id(db, request.form.get("subject_id"))
        try:
            minutes = int(request.form.get("minutes", app.config["LAB_KEY_DEFAULT_MINUTES"]))
        except ValueError:
            minutes = app.config["LAB_KEY_DEFAULT_MINUTES"]

        if not request.form.get("course_id"):
            flash("Select a course first.", "error")
            return redirect(url_for("admin_panel"))
        if course_id is None:
            flash("Select a valid course.", "error")
            return redirect(url_for("admin_panel"))
        if subject_id is not None and not db.execute(
            "SELECT 1 FROM course_subject WHERE course_id = ? AND subject_id = ?",
            (course_id, subject_id),
        ).fetchone():
            flash("Assign that subject to the selected course first.", "error")
            return redirect(url_for("admin_panel"))

        code = None
        for _ in range(10):
            candidate = f"{secrets.randbelow(900000) + 100000}"
            clash = db.execute(
                "SELECT id FROM lab_session WHERE code = ? AND is_active = 1", (candidate,)
            ).fetchone()
            if not clash:
                code = candidate
                break
        code = code or f"{secrets.randbelow(900000) + 100000}"

        expires_at = (
            datetime.now(timezone.utc).replace(tzinfo=None) + timedelta(minutes=minutes)
        ).strftime("%Y-%m-%d %H:%M:%S")
        db.execute(
            "INSERT INTO lab_session (course_id, subject_id, code, created_at, expires_at, is_active) "
            "VALUES (?, ?, ?, ?, ?, 1)",
            (course_id, subject_id, code, utcnow_iso(), expires_at),
        )
        db.commit()
        flash(f"Lab key generated: {code} (valid {minutes} min)", "success")
        return redirect(url_for("admin_panel"))

    @app.route("/admin/lab/<int:session_id>/deactivate", methods=["POST"])
    @login_required
    @admin_required
    def admin_deactivate_lab(session_id):
        db = get_db()
        db.execute("UPDATE lab_session SET is_active = 0 WHERE id = ?", (session_id,))
        db.commit()
        flash("Lab key deactivated.", "success")
        return redirect(url_for("admin_panel"))

    @app.route("/admin/attendance")
    @login_required
    @admin_required
    def admin_attendance():
        db = get_db()
        courses = db.execute("SELECT * FROM course ORDER BY name").fetchall()
        subjects = db.execute("SELECT * FROM subject ORDER BY name").fetchall()
        selected_course_id = request.args.get("course_id", "").strip()
        selected_section = request.args.get("section", "").strip().upper()
        if selected_section not in {"A", "B", "C", "D", "E", "F", "G"}:
            selected_section = ""
        selected_subject_id = request.args.get("subject_id", "").strip()
        selected_student_id = request.args.get("student_id", "").strip()
        current_ist = datetime.now(timezone(timedelta(hours=5, minutes=30)))
        selected_month = request.args.get("month", current_ist.strftime("%Y-%m"))
        try:
            month_date = datetime.strptime(selected_month, "%Y-%m")
        except ValueError:
            selected_month = current_ist.strftime("%Y-%m")
            month_date = datetime.strptime(selected_month, "%Y-%m")

        course_id = valid_course_id(db, selected_course_id)
        if course_id is None and selected_course_id:
            selected_course_id = ""
        subject_id = valid_subject_id(db, selected_subject_id)
        if subject_id is None and selected_subject_id:
            selected_subject_id = ""

        days_in_month = calendar.monthrange(month_date.year, month_date.month)[1]
        students = []
        attendance_by_student = {}
        lab_keys_by_student = {}
        student_query = "SELECT id, name, student_id, email, section FROM user WHERE role = 'student'"
        student_params = []
        attendance_query = (
            "SELECT attendance.student_id, attendance.timestamp, lab_session.code AS lab_code FROM attendance "
            "JOIN lab_session ON attendance.lab_session_id = lab_session.id"
        )
        attendance_params = []
        student_filters = []
        if course_id is not None:
            student_filters.append("course_id = ?")
            student_params.append(course_id)
            attendance_params.append(course_id)
        if subject_id is not None:
            student_filters.append(
                "course_id IN (SELECT course_id FROM course_subject WHERE subject_id = ?)"
            )
            student_params.append(subject_id)
            attendance_params.append(subject_id)
        if selected_section:
            student_filters.append("section = ?")
            student_params.append(selected_section)
        if selected_student_id:
            student_filters.append("student_id = ?")
            student_params.append(selected_student_id)
        if student_filters:
            student_query += " AND " + " AND ".join(student_filters)
        attendance_filters = []
        if course_id is not None:
            attendance_filters.append("lab_session.course_id = ?")
        if subject_id is not None:
            attendance_filters.append("lab_session.subject_id = ?")
        if attendance_filters:
            attendance_query += " WHERE " + " AND ".join(attendance_filters)
        student_query += " ORDER BY name"
        students = db.execute(student_query, student_params).fetchall()
        attendance_rows = db.execute(attendance_query, attendance_params).fetchall()
        for attendance in attendance_rows:
            try:
                utc_value = datetime.strptime(
                    attendance["timestamp"][:19], "%Y-%m-%d %H:%M:%S"
                ).replace(tzinfo=timezone.utc)
                ist_value = utc_value.astimezone(timezone(timedelta(hours=5, minutes=30)))
            except (TypeError, ValueError):
                continue
            if ist_value.strftime("%Y-%m") == selected_month:
                attendance_by_student.setdefault(attendance["student_id"], {})[
                    ist_value.day
                ] = ist_value.strftime("%H:%M")
                lab_keys_by_student.setdefault(attendance["student_id"], set()).add(
                    attendance["lab_code"]
                )

        matrix = []
        for student in students:
            marks = attendance_by_student.get(student["id"], {})
            present_count = len(marks)
            matrix.append({
                "student": student,
                "marks": marks,
                "lab_keys": sorted(lab_keys_by_student.get(student["id"], set())),
                "present_count": present_count,
                "absent_count": days_in_month - present_count,
                "percentage": round((present_count / days_in_month) * 100) if days_in_month else 0,
            })
        selected_course_name = ""
        if course_id is not None:
            selected_course = db.execute(
                "SELECT name, year FROM course WHERE id = ?",
                (course_id,),
            ).fetchone()
            if selected_course:
                selected_course_name = format_course_label(selected_course["name"], selected_course["year"])
        return render_template(
            "admin_attendance_overview.html",
            courses=courses,
            subjects=subjects,
            matrix=matrix,
            days=range(1, days_in_month + 1),
            selected_course_id=selected_course_id,
            selected_section=selected_section,
            selected_subject_id=selected_subject_id,
            selected_student_id=selected_student_id,
            selected_subject_name=(
                db.execute("SELECT name FROM subject WHERE id = ?", (subject_id,)).fetchone()["name"]
                if subject_id is not None else ""
            ),
            selected_month=selected_month,
            selected_course_name=selected_course_name,
        )

    @app.route("/admin/reports")
    @login_required
    @admin_required
    def admin_reports():
        db = get_db()
        date_from = request.args.get("date_from", "").strip()
        date_to = request.args.get("date_to", "").strip()
        student_id = request.args.get("student_id", "").strip()
        subject_id = request.args.get("subject_id", "").strip()
        selected_course_id = request.args.get("course_id", "").strip()
        course_id = valid_course_id(db, selected_course_id)
        courses = db.execute("SELECT * FROM course ORDER BY name, year").fetchall()
        report_query = (
            "SELECT attendance.timestamp, attendance.student_id, user.name AS student_name, "
            "user.student_id AS roll_no, user.email, user.section AS student_section, "
            "course.name AS course_name, course.year AS course_year, subject.name AS subject_name, lab_session.code "
            "FROM attendance "
            "JOIN user ON attendance.student_id = user.id "
            "JOIN lab_session ON attendance.lab_session_id = lab_session.id "
            "JOIN course ON lab_session.course_id = course.id "
            "LEFT JOIN subject ON lab_session.subject_id = subject.id"
        )
        report_filters = []
        report_params = []
        if course_id is not None:
            report_filters.append("course.id = ?")
            report_params.append(course_id)
        if date_from:
            report_filters.append("attendance.timestamp >= ?")
            report_params.append(f"{date_from} 00:00:00")
        if date_to:
            try:
                end_date = datetime.strptime(date_to, "%Y-%m-%d") + timedelta(days=1)
                report_filters.append("attendance.timestamp < ?")
                report_params.append(end_date.strftime("%Y-%m-%d 00:00:00"))
            except ValueError:
                date_to = ""
        if student_id:
            report_filters.append("user.student_id = ?")
            report_params.append(student_id)
        if subject_id.isdigit():
            report_filters.append("lab_session.subject_id = ?")
            report_params.append(int(subject_id))
        else:
            subject_id = ""
        if report_filters:
            report_query += " WHERE " + " AND ".join(report_filters)
        report_query += " ORDER BY attendance.timestamp DESC"
        records = db.execute(report_query, report_params).fetchall()
        subjects = db.execute("SELECT id, name FROM subject ORDER BY name").fetchall()
        return render_template(
            "admin_reports.html",
            records=records,
            courses=courses,
            subjects=subjects,
            date_from=date_from,
            date_to=date_to,
            selected_student_id=student_id,
            selected_subject_id=subject_id,
            selected_course_id=selected_course_id,
        )

    @app.route("/admin/attendance/export")
    @login_required
    @admin_required
    def admin_export_attendance():
        db = get_db()
        export_view = request.args.get("view", "attendance")
        course_id = request.args.get("course_id", "").strip()
        subject_id = request.args.get("subject_id", "").strip()
        month = request.args.get("month", "").strip()
        section = request.args.get("section", "").strip().upper()
        date_from = request.args.get("date_from", "").strip()
        date_to = request.args.get("date_to", "").strip()
        student_id = request.args.get("student_id", "").strip()

        query = (
            "SELECT attendance.student_id, attendance.timestamp, user.student_id AS roll_no, "
            "user.name AS student_name, user.email, "
            "course.name AS course_name, course.year AS course_year, user.section AS student_section, subject.name AS subject_name, "
            "lab_session.code "
            "FROM attendance JOIN user ON attendance.student_id = user.id "
            "JOIN lab_session ON attendance.lab_session_id = lab_session.id "
            "JOIN course ON lab_session.course_id = course.id "
            "LEFT JOIN subject ON lab_session.subject_id = subject.id"
        )
        filters = []
        params = []
        valid_course = valid_course_id(db, course_id)
        valid_subject = valid_subject_id(db, subject_id)
        if valid_course is not None:
            filters.append("lab_session.course_id = ?")
            params.append(valid_course)
        if valid_subject is not None:
            filters.append("lab_session.subject_id = ?")
            params.append(valid_subject)
        if student_id:
            filters.append("user.student_id = ?")
            params.append(student_id)
        if section in {"A", "B", "C", "D", "E", "F", "G"}:
            filters.append("user.section = ?")
            params.append(section)
        if export_view == "report":
            if date_from:
                filters.append("attendance.timestamp >= ?")
                params.append(f"{date_from} 00:00:00")
            if date_to:
                try:
                    end_date = datetime.strptime(date_to, "%Y-%m-%d") + timedelta(days=1)
                    filters.append("attendance.timestamp < ?")
                    params.append(end_date.strftime("%Y-%m-%d 00:00:00"))
                except ValueError:
                    pass
        if filters:
            query += " WHERE " + " AND ".join(filters)
        query += " ORDER BY attendance.timestamp DESC"
        records = db.execute(query, params).fetchall()

        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "Attendance"
        report_headers = [
            "Date & Time", "Student Name", "Email", "Roll No.", "Course", "Section", "Subject", "Lab Code",
        ]
        attendance_headers = [
            "S.No.", "Roll No.", "Student Name",
        ]
        if export_view == "report":
            worksheet.append(report_headers)
            for record in records:
                worksheet.append([
                    format_ist(record["timestamp"]),
                    record["student_name"],
                    record["email"],
                    record["roll_no"] or "",
                    format_course_label(record["course_name"], record["course_year"]),
                    record["student_section"] or "",
                    record["subject_name"] or "Legacy / Unassigned",
                    record["code"],
                ])
        else:
            try:
                month_date = datetime.strptime(month, "%Y-%m")
            except ValueError:
                month_date = datetime.now(timezone(timedelta(hours=5, minutes=30)))
                month = month_date.strftime("%Y-%m")
            days_in_month = calendar.monthrange(month_date.year, month_date.month)[1]
            student_query = "SELECT id, name, student_id, section FROM user WHERE role = 'student'"
            student_params = []
            student_filters = []
            if valid_course is not None:
                student_filters.append("course_id = ?")
                student_params.append(valid_course)
            if valid_subject is not None:
                student_filters.append(
                    "course_id IN (SELECT course_id FROM course_subject WHERE subject_id = ?)"
                )
                student_params.append(valid_subject)
            if student_id:
                student_filters.append("student_id = ?")
                student_params.append(student_id)
            if section in {"A", "B", "C", "D", "E", "F", "G"}:
                student_filters.append("section = ?")
                student_params.append(section)
            if student_filters:
                student_query += " AND " + " AND ".join(student_filters)
            students = db.execute(student_query + " ORDER BY name", student_params).fetchall()
            marks_by_student = {}
            for record in records:
                try:
                    utc_value = datetime.strptime(
                        record["timestamp"][:19], "%Y-%m-%d %H:%M:%S"
                    ).replace(tzinfo=timezone.utc)
                    ist_value = utc_value.astimezone(timezone(timedelta(hours=5, minutes=30)))
                except (TypeError, ValueError):
                    continue
                if ist_value.strftime("%Y-%m") == month:
                    marks_by_student.setdefault(record["student_id"], {})[
                        ist_value.day
                    ] = ist_value.strftime("%H:%M")

            worksheet.append(attendance_headers + ["Section"] + list(range(1, days_in_month + 1)) + [
                "Total Present", "Total Absent", "Attendance %",
            ])
            for index, student in enumerate(students, start=1):
                marks = marks_by_student.get(student["id"], {})
                present_count = len(marks)
                worksheet.append([
                    index,
                    student["student_id"] or "-",
                    student["name"],
                    student["section"] or "-",
                    *[marks.get(day, "A") for day in range(1, days_in_month + 1)],
                    present_count,
                    days_in_month - present_count,
                    f"{round((present_count / days_in_month) * 100) if days_in_month else 0}%",
                ])
        for cell in worksheet[1]:
            cell.font = openpyxl.styles.Font(bold=True)
        widths = [24, 24, 32, 16, 28, 12, 24, 14] if export_view == "report" else [8, 16, 24, 12] + [10] * days_in_month + [16, 16, 16]
        for index, width in enumerate(widths, start=1):
            worksheet.column_dimensions[openpyxl.utils.get_column_letter(index)].width = width
        output = io.BytesIO()
        workbook.save(output)
        output.seek(0)
        return send_file(
            output,
            as_attachment=True,
            download_name="reports.xlsx" if export_view == "report" else "attendance.xlsx",
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.route("/admin/lab/<int:session_id>/attendance")
    @login_required
    @admin_required
    def admin_view_attendance(session_id):
        db = get_db()
        session_row = db.execute(
            "SELECT lab_session.*, course.name AS course_name, course.year AS course_year FROM lab_session "
            "JOIN course ON lab_session.course_id = course.id WHERE lab_session.id = ?",
            (session_id,),
        ).fetchone()
        if not session_row:
            abort(404)
        records = db.execute(
            "SELECT attendance.*, user.name AS student_name, user.email AS student_email, user.section AS student_section "
            "FROM attendance JOIN user ON attendance.student_id = user.id "
            "WHERE attendance.lab_session_id = ? ORDER BY attendance.timestamp",
            (session_id,),
        ).fetchall()
        return render_template("admin_attendance.html", session=session_row, records=records)

    @app.errorhandler(403)
    def forbidden(e):
        return render_template("403.html"), 403

    @app.errorhandler(404)
    def page_not_found(e):
        return render_template("404.html"), 404

    @app.errorhandler(500)
    def server_error(e):
        return render_template("500.html"), 500


if __name__ == "__main__":
    app = create_app()
    app.run(host="0.0.0.0", port=5000, debug=True, use_reloader=False)